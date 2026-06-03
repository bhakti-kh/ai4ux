# ============================================================
# Ai4UX — Sprint 7C: RAG + Figma Push + Team Features
# ============================================================
from flask import Flask, request, render_template_string, jsonify, send_file, send_from_directory, session, redirect
import anthropic, requests as http_requests, json, io, base64 as b64_lib
import sqlite3, os
from datetime import datetime
from functools import wraps
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib import colors
import urllib.parse, secrets as _secrets_mod

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "ai4ux-dev-secret-change-in-production")
app.config['MAX_CONTENT_LENGTH'] = 50 * 1024 * 1024

DIST_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "dist")
USE_REACT = os.path.exists(DIST_DIR)

# ============================================================
# CREDENTIALS
# ============================================================
FIGMA_TOKEN    = os.environ.get("FIGMA_TOKEN", "")
FIGMA_FILE_KEY = os.environ.get("FIGMA_FILE_KEY", "")
FIGMA_API_BASE = "https://api.figma.com/v1"
ADMIN_EMAIL    = "bhakti.khandekar@gmail.com"
CLAUDE_API_KEY       = os.environ.get("CLAUDE_API_KEY", "")
JIRA_DOMAIN          = os.environ.get("JIRA_DOMAIN", "yourcompany.atlassian.net")
JIRA_EMAIL           = os.environ.get("JIRA_EMAIL", "you@company.com")
JIRA_API_TOKEN       = os.environ.get("JIRA_API_TOKEN", "")
GOOGLE_CLIENT_ID     = os.environ.get("GOOGLE_CLIENT_ID", "")
GOOGLE_CLIENT_SECRET = os.environ.get("GOOGLE_CLIENT_SECRET", "")
GOOGLE_REDIRECT_URI  = os.environ.get("GOOGLE_REDIRECT_URI", "http://localhost:5000/auth/callback")
GOOGLE_AUTH_URL      = "https://accounts.google.com/o/oauth2/v2/auth"
GOOGLE_TOKEN_URL     = "https://oauth2.googleapis.com/token"
GOOGLE_USERINFO_URL  = "https://www.googleapis.com/oauth2/v2/userinfo"

# ============================================================
# DATABASE
# ============================================================
DATABASE_URL = os.environ.get("DATABASE_URL", "")

if DATABASE_URL:
    try:
        import psycopg2, psycopg2.extras
        _USE_PG = True
        print("Using PostgreSQL")
    except ImportError:
        _USE_PG = False
        print("psycopg2 not found, falling back to SQLite")
else:
    _USE_PG = False
    print("Using SQLite (local dev)")

class _PGCursor:
    def __init__(self, cur): self._cur = cur
    def fetchall(self): return [dict(r) for r in (self._cur.fetchall() or [])]
    def fetchone(self):
        r = self._cur.fetchone()
        return dict(r) if r else None
    @property
    def lastrowid(self):
        try: self._cur.execute("SELECT lastval()"); return self._cur.fetchone()[0]
        except: return None

class _PGConn:
    def __init__(self, dsn):
        self._conn = psycopg2.connect(dsn)
        self.row_factory = None
    def execute(self, query, params=()):
        query = (query
            .replace("?", "%s")
            .replace("INTEGER PRIMARY KEY AUTOINCREMENT", "SERIAL PRIMARY KEY")
            .replace("INTEGER PRIMARY KEY AUTOINCREMENT,", "SERIAL PRIMARY KEY,"))
        cur = self._conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
        cur.execute(query, params or ())
        return _PGCursor(cur)
    def commit(self):  self._conn.commit()
    def close(self):   self._conn.close()
    def __enter__(self): return self
    def __exit__(self, *a): self.close()

def connect_db():
    if _USE_PG:
        return _PGConn(DATABASE_URL)
    import sqlite3 as _sq
    _conn = _sq.connect("ai4ux.db")
    _conn.row_factory = _sq.Row
    return _conn

client        = anthropic.Anthropic(api_key=CLAUDE_API_KEY)
last_analysis = {}

# ============================================================
# RAG
# ============================================================
try:
    from rag_store import rag
    _RAG_AVAILABLE = True
except Exception as _rag_err:
    _RAG_AVAILABLE = False
    print(f"RAG not available: {_rag_err}")
    class _RagStub:
        def ingest_all(self, *a, **kw): pass
        def ingest_convention(self, *a, **kw): pass
        def ingest_component(self, *a, **kw): pass
        def ingest_ds_file(self, *a, **kw): pass
        def ingest_custom_guidelines(self, *a, **kw): pass
        def ingest_product_context(self, *a, **kw): pass
        def ingest_builtin_guidelines(self, *a, **kw): pass
        def build_analysis_context(self, *a, **kw): return ""
        def build_component_gen_context(self, *a, **kw): return ""
        def stats(self, *a, **kw): return {"available": False}
    rag = _RagStub()

# Ingest built-in guidelines at startup
try:
    pass  # will be called after BUILTIN_GUIDELINES is defined
except: pass

def _rag_warmup_async(user_id):
    import threading
    def _do():
        try:
            with connect_db() as conn:
                rag.ingest_all(conn, user_id, builtin_guidelines=BUILTIN_GUIDELINES)
        except Exception as e:
            app.logger.warning(f"RAG warmup failed for {user_id}: {e}")
    threading.Thread(target=_do, daemon=True).start()


# ============================================================
# GUIDELINES REGISTRY
# ============================================================
BUILTIN_GUIDELINES = {
  "wcag_aa": {
    "name": "WCAG 2.2 AA", "version": "2.2", "category": "accessibility",
    "description": "Web Content Accessibility Guidelines Level AA",
    "url": "https://www.w3.org/TR/WCAG22/",
    "criteria": {
      "1.1.1": {"title":"Non-text Content","level":"A","components":["Image","Icon","Button","Avatar","Chart"],"description":"All non-text content has a text alternative","url":"https://www.w3.org/WAI/WCAG22/Understanding/non-text-content.html","fix":"Add alt text to images, aria-label to icon buttons, title to charts"},
      "1.3.1": {"title":"Info and Relationships","level":"A","components":["Form","Table","List","Heading","Input"],"description":"Information conveyed through presentation can be programmatically determined","url":"https://www.w3.org/WAI/WCAG22/Understanding/info-and-relationships.html","fix":"Use semantic HTML: label elements for inputs, th for headers, fieldset/legend for groups"},
      "1.4.1": {"title":"Use of Color","level":"A","components":["Chart","Status","Alert","Tag","Badge","Link"],"description":"Color is not the only visual means of conveying information","url":"https://www.w3.org/WAI/WCAG22/Understanding/use-of-color.html","fix":"Add icons, patterns, or text labels alongside color coding"},
      "1.4.3": {"title":"Contrast (Minimum)","level":"AA","components":["Button","Input","Text","Label","Tag","Link","Badge"],"description":"Text has contrast ratio of at least 4.5:1","url":"https://www.w3.org/WAI/WCAG22/Understanding/contrast-minimum.html","fix":"Use a contrast checker. Minimum: 4.5:1 normal text, 3:1 large text"},
      "1.4.4": {"title":"Resize Text","level":"AA","components":["Text","Label","Input","Card"],"description":"Text can be resized up to 200% without loss of content","url":"https://www.w3.org/WAI/WCAG22/Understanding/resize-text.html","fix":"Use relative units (rem/em) for font sizes"},
      "1.4.10": {"title":"Reflow","level":"AA","components":["Layout","Container","Navigation","Table"],"description":"Content can reflow to single column at 320px width","url":"https://www.w3.org/WAI/WCAG22/Understanding/reflow.html","fix":"Use responsive CSS, avoid fixed-width layouts"},
      "1.4.11": {"title":"Non-text Contrast","level":"AA","components":["Input","Button","Checkbox","Radio","Chart","Icon"],"description":"UI components have 3:1 contrast against adjacent colors","url":"https://www.w3.org/WAI/WCAG22/Understanding/non-text-contrast.html","fix":"Ensure input borders, focus rings, and UI icons meet 3:1 against background"},
      "1.4.12": {"title":"Text Spacing","level":"AA","components":["Text","Input","Card","Label"],"description":"No loss of content when text spacing is overridden","url":"https://www.w3.org/WAI/WCAG22/Understanding/text-spacing.html","fix":"Test with text spacing bookmarklet; avoid overflow:hidden that clips expanded text"},
      "2.1.1": {"title":"Keyboard","level":"A","components":["Button","Input","Dropdown","Modal","Tab","Navigation","Link","Checkbox"],"description":"All functionality is available from a keyboard","url":"https://www.w3.org/WAI/WCAG22/Understanding/keyboard.html","fix":"Ensure all interactive elements are focusable and operable with Enter/Space/Arrow keys"},
      "2.1.2": {"title":"No Keyboard Trap","level":"A","components":["Modal","Dropdown","Overlay","Dialog"],"description":"Keyboard focus can always be moved away from any component","url":"https://www.w3.org/WAI/WCAG22/Understanding/no-keyboard-trap.html","fix":"Implement proper focus trap in modals with Escape key to close"},
      "2.4.3": {"title":"Focus Order","level":"A","components":["Form","Modal","Navigation","Wizard"],"description":"Focus order preserves meaning and operability","url":"https://www.w3.org/WAI/WCAG22/Understanding/focus-order.html","fix":"Ensure DOM order matches visual order; avoid tabindex > 0"},
      "2.4.6": {"title":"Headings and Labels","level":"AA","components":["Form","Table","Section","Input","Page"],"description":"Headings and labels describe topic or purpose","url":"https://www.w3.org/WAI/WCAG22/Understanding/headings-and-labels.html","fix":"Add descriptive aria-label or label elements; use heading hierarchy (h1-h6) appropriately"},
      "2.4.7": {"title":"Focus Visible","level":"AA","components":["Button","Input","Link","Tab","Checkbox","Radio"],"description":"Any keyboard operable interface has a visible keyboard focus indicator","url":"https://www.w3.org/WAI/WCAG22/Understanding/focus-visible.html","fix":"Never use outline:none without a replacement"},
      "2.4.11": {"title":"Focus Not Obscured (Minimum)","level":"AA","components":["Navigation","Sticky Header","Modal","Tooltip"],"description":"Focused component is not entirely hidden by author-created content","url":"https://www.w3.org/WAI/WCAG22/Understanding/focus-not-obscured-minimum.html","fix":"Ensure sticky headers/footers don't completely cover focused elements"},
      "2.5.3": {"title":"Label in Name","level":"A","components":["Button","Input","Link","Icon Button"],"description":"For components with visible text labels, accessible name contains the visible text","url":"https://www.w3.org/WAI/WCAG22/Understanding/label-in-name.html","fix":"Ensure aria-label starts with or contains the visible button/link text"},
      "3.3.1": {"title":"Error Identification","level":"A","components":["Form","Input","Validation"],"description":"Input errors are automatically detected and described to the user in text","url":"https://www.w3.org/WAI/WCAG22/Understanding/error-identification.html","fix":"Show error messages as text near the field; use aria-invalid=true"},
      "3.3.2": {"title":"Labels or Instructions","level":"A","components":["Form","Input","Field"],"description":"Labels or instructions are provided when content requires user input","url":"https://www.w3.org/WAI/WCAG22/Understanding/labels-or-instructions.html","fix":"Add visible labels (not just placeholder text)"},
      "4.1.2": {"title":"Name, Role, Value","level":"A","components":["Button","Input","Checkbox","Select","Modal","Tab","Toggle"],"description":"All UI components have name, role, and value that can be programmatically determined","url":"https://www.w3.org/WAI/WCAG22/Understanding/name-role-value.html","fix":"Use semantic HTML or ARIA roles; ensure aria-expanded, aria-selected, aria-checked are set correctly"},
      "4.1.3": {"title":"Status Messages","level":"AA","components":["Alert","Notification","Toast","Loading","Success"],"description":"Status messages can be programmatically determined","url":"https://www.w3.org/WAI/WCAG22/Understanding/status-messages.html","fix":"Use role=alert for urgent messages, role=status for non-urgent; use aria-live regions"}
    }
  },
  "nielsen": {
    "name": "Nielsen Norman 10 Heuristics", "version": "1994", "category": "ux",
    "description": "Jakob Nielsen's 10 general principles for interaction design",
    "url": "https://www.nngroup.com/articles/ten-usability-heuristics/",
    "criteria": {
      "H1": {"title":"Visibility of System Status","components":["Loading","Progress","Status","Notification","Toast"],"description":"Always keep users informed about what is going on","url":"https://www.nngroup.com/articles/visibility-system-status/","fix":"Add loading indicators, progress bars, success/error confirmations for all user actions"},
      "H2": {"title":"Match Between System and Real World","components":["Navigation","Label","Button","Form","Icon"],"description":"Use words, phrases and concepts familiar to the user","url":"https://www.nngroup.com/articles/match-system-real-world/","fix":"Avoid technical jargon; use task-oriented language"},
      "H3": {"title":"User Control and Freedom","components":["Modal","Form","Navigation","Wizard","Delete"],"description":"Support undo and redo; let users exit unwanted states","url":"https://www.nngroup.com/articles/user-control-and-freedom/","fix":"Add Cancel buttons, Undo functionality, clear Exit options, Back navigation"},
      "H4": {"title":"Consistency and Standards","components":["Button","Navigation","Form","Icon","Layout"],"description":"Follow platform conventions","url":"https://www.nngroup.com/articles/consistency-and-standards/","fix":"Use consistent terminology, placement, and visual treatment for same-function elements"},
      "H5": {"title":"Error Prevention","components":["Form","Input","Confirmation","Delete","Submit"],"description":"Design to prevent problems from occurring","url":"https://www.nngroup.com/articles/slips/","fix":"Add input validation, confirmation dialogs, constraints on input fields"},
      "H6": {"title":"Recognition Rather Than Recall","components":["Navigation","Menu","Form","Label","Search","Dropdown"],"description":"Minimise memory load by making options, actions and objects visible","url":"https://www.nngroup.com/articles/recognition-and-recall/","fix":"Show options rather than requiring recall; use autocomplete, recent searches"},
      "H7": {"title":"Flexibility and Efficiency of Use","components":["Navigation","Search","Form","Table","Keyboard"],"description":"Accelerators for expert users; allow users to tailor frequent actions","url":"https://www.nngroup.com/articles/flexibility-efficiency-of-use/","fix":"Add keyboard shortcuts, saved filters, bulk actions"},
      "H8": {"title":"Aesthetic and Minimalist Design","components":["Dashboard","Card","Modal","Form","Page"],"description":"Avoid irrelevant information","url":"https://www.nngroup.com/articles/aesthetic-and-minimalist-design/","fix":"Remove decorative elements that don't aid comprehension; reduce visual noise"},
      "H9": {"title":"Help Users Recognize, Diagnose and Recover From Errors","components":["Form","Alert","Input","Validation"],"description":"Error messages in plain language, precisely indicating the problem","url":"https://www.nngroup.com/articles/error-message-guidelines/","fix":"Write human-readable errors with specific cause and actionable fix"},
      "H10": {"title":"Help and Documentation","components":["Tooltip","Help","Empty State","Onboarding"],"description":"Provide easy-to-search help documentation focused on user tasks","url":"https://www.nngroup.com/articles/help-and-documentation/","fix":"Add contextual tooltips, inline help text, empty state guidance"}
    }
  },
  "wai_aria": {
    "name": "WAI-ARIA Authoring Practices", "version": "1.2", "category": "accessibility",
    "description": "W3C patterns for implementing accessible rich internet applications",
    "url": "https://www.w3.org/WAI/ARIA/apg/patterns/",
    "criteria": {
      "ARIA-BTN":  {"title":"Button Pattern","components":["Button","Icon Button","Toggle Button"],"description":"role=button, Enter/Space to activate, aria-pressed for toggle","url":"https://www.w3.org/WAI/ARIA/apg/patterns/button/","fix":"Add role=button if not <button>, aria-label for icon-only, aria-pressed for toggles"},
      "ARIA-DLG":  {"title":"Dialog (Modal) Pattern","components":["Modal","Dialog","Overlay"],"description":"role=dialog, aria-modal=true, focus trap inside, Escape to close","url":"https://www.w3.org/WAI/ARIA/apg/patterns/dialog-modal/","fix":"Implement focus trap, add aria-modal=true, aria-labelledby, restore focus when closed"},
      "ARIA-FORM": {"title":"Form Pattern","components":["Form","Input","Field Group"],"description":"aria-label on form, aria-required, aria-invalid, aria-describedby for errors","url":"https://www.w3.org/WAI/ARIA/apg/patterns/","fix":"Link labels to inputs, add aria-required=true, set aria-invalid=true on error"},
      "ARIA-NAV":  {"title":"Navigation Landmark","components":["Navigation","Menu","Sidebar"],"description":"role=navigation, aria-label when multiple navs exist, skip link","url":"https://www.w3.org/WAI/ARIA/apg/patterns/","fix":"Use <nav> element, add aria-label for multiple navs"},
      "ARIA-TABS": {"title":"Tabs Pattern","components":["Tab","Tab Panel"],"description":"role=tablist, role=tab, role=tabpanel, aria-selected, arrow key navigation","url":"https://www.w3.org/WAI/ARIA/apg/patterns/tabs/","fix":"Add proper roles, aria-selected=true on active tab"},
      "ARIA-ALRT": {"title":"Alert and Live Region Pattern","components":["Alert","Notification","Toast","Banner","Status"],"description":"role=alert for urgent, role=status for non-urgent","url":"https://www.w3.org/WAI/ARIA/apg/patterns/alert/","fix":"Use role=alert for errors/urgent, aria-live=polite for status updates"},
      "ARIA-CHK":  {"title":"Checkbox Pattern","components":["Checkbox","Toggle","Switch"],"description":"role=checkbox, aria-checked, Space to toggle","url":"https://www.w3.org/WAI/ARIA/apg/patterns/checkbox/","fix":"Use <input type=checkbox> or role=checkbox with aria-checked"},
      "ARIA-CMBO": {"title":"Combobox Pattern","components":["Dropdown","Select","Autocomplete","Search"],"description":"role=combobox, aria-expanded, aria-autocomplete, arrow keys to navigate","url":"https://www.w3.org/WAI/ARIA/apg/patterns/combobox/","fix":"Implement role=combobox with aria-expanded, manage aria-activedescendant"}
    }
  },
  "carbon_a11y": {
    "name": "Carbon Design System Accessibility", "version": "11", "category": "platform",
    "description": "IBM Carbon Design System accessibility guidelines",
    "url": "https://carbondesignsystem.com/guidelines/accessibility/overview/",
    "criteria": {
      "CBN-BTN":   {"title":"Carbon Button Accessibility","components":["Button"],"description":"Use Carbon Button; icon buttons need aria-label","url":"https://carbondesignsystem.com/components/button/accessibility/","fix":"Always use cds-button; add hasIconOnly + iconDescription for icon-only buttons"},
      "CBN-FORM":  {"title":"Carbon Form Accessibility","components":["Form","Input","TextInput","Select"],"description":"Use Carbon FormItem wrapper; never use placeholder as label","url":"https://carbondesignsystem.com/components/text-input/accessibility/","fix":"Wrap all inputs in FormItem; add labelText prop; use helperText for format hints"},
      "CBN-COLOR": {"title":"Carbon Color Tokens","components":["All"],"description":"Use Carbon color tokens; never hardcode hex","url":"https://carbondesignsystem.com/guidelines/color/tokens/","fix":"Replace hardcoded colors with Carbon tokens: $blue-60 for interactive"},
      "CBN-FOCUS": {"title":"Carbon Focus Styling","components":["All interactive"],"description":"Never suppress Carbon's default focus ring","url":"https://carbondesignsystem.com/guidelines/accessibility/keyboard/","fix":"Remove outline:none overrides; use $focus (#0f62fe) 2px solid offset 1px"},
      "CBN-TYPE":  {"title":"Carbon Typography Scale","components":["Text","Heading","Label"],"description":"Use Carbon type tokens; body-01 for regular text","url":"https://carbondesignsystem.com/guidelines/typography/type-sets/","fix":"Use Carbon type tokens; IBM Plex Sans/Mono only; min 14px body text"}
    }
  },
  "material_a11y": {
    "name": "Material Design Accessibility", "version": "3", "category": "platform",
    "description": "Google Material Design 3 accessibility guidelines",
    "url": "https://m3.material.io/foundations/accessible-design/accessibility-basics",
    "criteria": {
      "MAT-TARGET": {"title":"Touch Target Size","components":["Button","Icon Button","Checkbox","Radio","FAB"],"description":"Minimum 48x48dp touch targets","url":"https://m3.material.io/foundations/accessible-design/accessibility-basics","fix":"Ensure all tappable elements meet 48x48dp minimum"},
      "MAT-COLOR":  {"title":"Dynamic Color Accessibility","components":["All","Theme","Button","Card"],"description":"Dynamic color must maintain WCAG AA contrast","url":"https://m3.material.io/styles/color/the-color-system/color-roles","fix":"Test dynamic color combinations with contrast checker"},
      "MAT-MOTION": {"title":"Reduce Motion","components":["Animation","Transition","Loading"],"description":"Respect prefers-reduced-motion","url":"https://m3.material.io/foundations/accessible-design/accessibility-basics","fix":"Add @media (prefers-reduced-motion: reduce) to disable or simplify all animations"},
      "MAT-FORM":   {"title":"Material Form Accessibility","components":["TextField","Outlined Input","Filled Input"],"description":"Use Supporting Text for requirements; never placeholder-only","url":"https://m3.material.io/components/text-fields/accessibility","fix":"Add supporting text for hints, error text with role=alert for validation"}
    }
  }
}

KNOWN_CONFLICTS = [
  {"id":"CF-001","title":"Ghost/Outlined Buttons","registry_a":"carbon_a11y","criterion_a":"CBN-BTN","registry_b":"wcag_aa","criterion_b":"1.4.11","description":"Carbon ghost buttons often fail WCAG 1.4.11 non-text contrast on light backgrounds.","interpretation_a":"Use Carbon ghost buttons as designed — acceptable per Carbon guidelines","interpretation_b":"Ghost buttons violate WCAG 1.4.11 — add a visible border of sufficient contrast"},
  {"id":"CF-002","title":"Placeholder Text as Labels","registry_a":"nielsen","criterion_a":"H6","registry_b":"wcag_aa","criterion_b":"3.3.2","description":"Placeholder-only inputs fail WCAG 3.3.2 as placeholder disappears on input.","interpretation_a":"Placeholder text provides sufficient recognition cue for simple search inputs","interpretation_b":"WCAG 3.3.2 requires always-visible labels — placeholder alone fails"},
  {"id":"CF-003","title":"Icon-Only Buttons","registry_a":"nielsen","criterion_a":"H8","registry_b":"wcag_aa","criterion_b":"1.1.1","description":"Minimalist icon-only buttons conflict with WCAG 1.1.1 text alternatives requirement.","interpretation_a":"Icon-only buttons reduce clutter — acceptable if icons are universally understood","interpretation_b":"All icon buttons must have aria-label or visually hidden text"},
  {"id":"CF-005","title":"Color-Only Status Indicators","registry_a":"nielsen","criterion_a":"H1","registry_b":"wcag_aa","criterion_b":"1.4.1","description":"Color alone for status violates WCAG 1.4.1.","interpretation_a":"Color-only status is acceptable in expert dashboards where users learn the colour coding","interpretation_b":"WCAG 1.4.1 requires a secondary indicator — add icons, labels, or patterns alongside colour coding"}
]


# ============================================================
# GUIDELINES DB FUNCTIONS
# ============================================================
def init_guidelines_db():
    conn = connect_db()
    conn.execute("""CREATE TABLE IF NOT EXISTS guideline_activation (
        id SERIAL PRIMARY KEY, user_id TEXT DEFAULT 'anonymous',
        registry_id TEXT, is_active INTEGER DEFAULT 1,
        scope TEXT DEFAULT 'global', product_name TEXT DEFAULT '', created_at TEXT)""")
    conn.execute("""CREATE TABLE IF NOT EXISTS custom_guidelines (
        id SERIAL PRIMARY KEY, user_id TEXT DEFAULT 'anonymous',
        name TEXT, version TEXT, category TEXT, content TEXT, filename TEXT,
        is_active INTEGER DEFAULT 1, created_at TEXT)""")
    conn.execute("""CREATE TABLE IF NOT EXISTS guideline_profiles (
        id SERIAL PRIMARY KEY, user_id TEXT DEFAULT 'anonymous',
        product_name TEXT, active_registries TEXT, overrides TEXT, created_at TEXT)""")
    conn.execute("""CREATE TABLE IF NOT EXISTS conflict_resolutions (
        id SERIAL PRIMARY KEY, user_id TEXT DEFAULT 'anonymous',
        conflict_id TEXT, resolution TEXT, remember INTEGER DEFAULT 0, created_at TEXT)""")
    conn.commit(); conn.close()

def get_active_registries(user_id="anonymous", product_name=None):
    conn = connect_db()
    rows = conn.execute(
        "SELECT registry_id, is_active, scope, product_name FROM guideline_activation WHERE user_id=?",
        (user_id,)).fetchall()
    conn.close()
    activations = {r['registry_id']: dict(r) for r in rows}
    result = []
    for rid in BUILTIN_GUIDELINES:
        if rid in activations:
            act = activations[rid]
            if act['scope'] == 'global' or act['product_name'] == (product_name or ''):
                if act['is_active']: result.append(rid)
        else:
            result.append(rid)
    return result

def set_registry_active(user_id, registry_id, is_active, scope='global', product_name=''):
    conn = connect_db()
    conn.execute("DELETE FROM guideline_activation WHERE user_id=? AND registry_id=? AND scope=? AND product_name=?",
        (user_id, registry_id, scope, product_name))
    conn.execute("INSERT INTO guideline_activation (user_id,registry_id,is_active,scope,product_name,created_at) VALUES (?,?,?,?,?,?)",
        (user_id, registry_id, 1 if is_active else 0, scope, product_name, datetime.now().strftime("%Y-%m-%d %H:%M")))
    conn.commit(); conn.close()

def get_conflict_resolution(user_id, conflict_id):
    conn = connect_db()
    row = conn.execute(
        "SELECT * FROM conflict_resolutions WHERE user_id=? AND conflict_id=? ORDER BY id DESC LIMIT 1",
        (user_id, conflict_id)).fetchone()
    conn.close()
    return dict(row) if row else None

def save_conflict_resolution(user_id, conflict_id, resolution, remember=False):
    conn = connect_db()
    if remember:
        conn.execute("DELETE FROM conflict_resolutions WHERE user_id=? AND conflict_id=?", (user_id, conflict_id))
    conn.execute("INSERT INTO conflict_resolutions (user_id,conflict_id,resolution,remember,created_at) VALUES (?,?,?,?,?)",
        (user_id, conflict_id, resolution, 1 if remember else 0, datetime.now().strftime("%Y-%m-%d %H:%M")))
    conn.commit(); conn.close()

def get_guidelines_context(user_id="anonymous", product_name=None, component_types=None):
    active_ids = get_active_registries(user_id, product_name)
    if not active_ids: return ""
    ctx = "ACTIVE DESIGN GUIDELINES (cite these in your analysis):\n"
    for rid in active_ids:
        if rid not in BUILTIN_GUIDELINES: continue
        reg = BUILTIN_GUIDELINES[rid]
        ctx += f"\n[{reg['name']}]\n"
        for cid, item in reg['criteria'].items():
            if component_types:
                item_comps = [c.lower() for c in item.get('components',[])]
                if not any(ct.lower() in ' '.join(item_comps) or any(ct.lower() in c for c in item_comps) for ct in component_types):
                    if 'all' not in item_comps: continue
            ctx += f"  {cid}: {item['title']} — {item['description'][:120]}\n"
    conn = connect_db()
    custom = conn.execute("SELECT * FROM custom_guidelines WHERE user_id=? AND is_active=1",(user_id,)).fetchall()
    conn.close()
    for c in custom:
        ctx += f"\n[{c['name']}]\n{c['content'][:1000]}\n"
    ctx += "\nFor each gap and recommendation, cite the specific criterion ID, explain the violation, provide the fix, and include the reference URL.\n"
    return ctx[:6000]

def get_active_conflicts(user_id="anonymous"):
    active_ids = set(get_active_registries(user_id))
    relevant = []
    for conflict in KNOWN_CONFLICTS:
        if conflict['registry_a'] in active_ids and conflict['registry_b'] in active_ids:
            resolution = get_conflict_resolution(user_id, conflict['id'])
            relevant.append({**conflict, 'resolution': resolution})
    return relevant

# ============================================================
# TEAM DB INIT
# ============================================================
def init_teams_db():
    conn = connect_db()
    conn.execute("""CREATE TABLE IF NOT EXISTS teams (
        id SERIAL PRIMARY KEY, name TEXT NOT NULL,
        created_by TEXT NOT NULL, invite_code TEXT UNIQUE NOT NULL, created_at TEXT)""")
    conn.execute("""CREATE TABLE IF NOT EXISTS team_members (
        id SERIAL PRIMARY KEY, team_id INTEGER NOT NULL,
        user_email TEXT NOT NULL, role TEXT DEFAULT 'member',
        status TEXT DEFAULT 'pending', invited_by TEXT, joined_at TEXT,
        UNIQUE(team_id, user_email))""")
    conn.execute("""CREATE TABLE IF NOT EXISTS team_invites (
        id SERIAL PRIMARY KEY, team_id INTEGER NOT NULL,
        invited_email TEXT NOT NULL, invited_by TEXT NOT NULL,
        token TEXT UNIQUE NOT NULL, status TEXT DEFAULT 'pending', created_at TEXT)""")
    conn.execute("""CREATE TABLE IF NOT EXISTS spec_sharing (
        id SERIAL PRIMARY KEY, spec_id TEXT NOT NULL,
        shared_by TEXT NOT NULL, team_id INTEGER NOT NULL, created_at TEXT,
        UNIQUE(spec_id, team_id))""")
    conn.commit(); conn.close()

# ============================================================
# DB PATHS + AUTH HELPERS
# ============================================================
SPECS_DIR = "specs"; DS_DIR = "design_system"; RESEARCH_DIR = "research"; DB_PATH = "ai4ux.db"

def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if 'user' not in session: return redirect('/login')
        return f(*args, **kwargs)
    return decorated

def current_user():    return session.get('user', {})
def current_user_id(): return session.get('user', {}).get('email', 'anonymous')

# ============================================================
# DB INIT CALLS
# ============================================================
def init_db():
    os.makedirs(SPECS_DIR, exist_ok=True)
    conn = connect_db()
    conn.execute("""CREATE TABLE IF NOT EXISTS specs (
        id SERIAL PRIMARY KEY, spec_id TEXT UNIQUE, ticket_id TEXT,
        ticket_summary TEXT, screen_file TEXT, compliance INTEGER, components TEXT,
        created_at TEXT, json_path TEXT, user_id TEXT DEFAULT 'anonymous', spec_data TEXT)""")
    conn.commit(); conn.close()

def init_ds_db():
    os.makedirs(DS_DIR, exist_ok=True)
    conn = connect_db()
    conn.execute("""CREATE TABLE IF NOT EXISTS ds_files (
        id SERIAL PRIMARY KEY, filename TEXT, file_type TEXT,
        content TEXT, uploaded_at TEXT, user_id TEXT DEFAULT 'anonymous')""")
    conn.commit(); conn.close()

def init_research_db():
    os.makedirs(RESEARCH_DIR, exist_ok=True)
    conn = connect_db()
    conn.execute("""CREATE TABLE IF NOT EXISTS research_items (
        id SERIAL PRIMARY KEY, title TEXT, item_type TEXT,
        source TEXT, content TEXT, added_at TEXT, user_id TEXT DEFAULT 'anonymous')""")
    conn.commit(); conn.close()

def init_conventions_db():
    conn = connect_db()
    conn.execute("""CREATE TABLE IF NOT EXISTS conventions (
        id SERIAL PRIMARY KEY, user_id TEXT DEFAULT 'anonymous',
        title TEXT, description TEXT, category TEXT, priority TEXT,
        source_ticket TEXT, source_screen TEXT, feedback_type TEXT, created_at TEXT)""")
    conn.commit(); conn.close()

def init_generated_components_db():
    conn = connect_db()
    conn.execute("""CREATE TABLE IF NOT EXISTS generated_components (
        id SERIAL PRIMARY KEY, user_id TEXT DEFAULT 'anonymous',
        name TEXT, type TEXT, html_code TEXT, react_code TEXT, design_specs TEXT,
        source_ticket TEXT, source_screen TEXT,
        status TEXT DEFAULT 'canonical', conflict_with INTEGER, created_at TEXT)""")
    conn.commit(); conn.close()

def init_product_context_db():
    conn = connect_db()
    conn.execute("""CREATE TABLE IF NOT EXISTS product_context (
        id SERIAL PRIMARY KEY, user_id TEXT DEFAULT 'anonymous',
        product_name TEXT DEFAULT 'Unknown', feature_domain TEXT, pattern TEXT, insight TEXT,
        source_ticket TEXT, source_screen TEXT,
        item_type TEXT DEFAULT 'auto', content TEXT, title TEXT, created_at TEXT)""")
    conn.commit(); conn.close()

def init_gap_resolutions_db():
    conn = connect_db()
    conn.execute("""CREATE TABLE IF NOT EXISTS gap_resolutions (
        id SERIAL PRIMARY KEY, user_id TEXT DEFAULT 'anonymous',
        spec_id TEXT, gap_title TEXT, status TEXT DEFAULT 'open',
        notes TEXT, created_at TEXT)""")
    conn.commit(); conn.close()

def init_init_db_route_table():
    """Emergency init — hit /init-db-now to create all tables."""
    pass

init_db()
init_ds_db()
init_research_db()
init_conventions_db()
init_generated_components_db()
init_product_context_db()
init_guidelines_db()
init_gap_resolutions_db()
init_teams_db()

# Startup: ingest built-in guidelines into RAG
try:
    rag.ingest_builtin_guidelines(BUILTIN_GUIDELINES)
except Exception as _e:
    print(f"Startup guideline ingest skipped: {_e}")


# ============================================================
# TEAM HELPER FUNCTIONS
# ============================================================
def get_user_team(user_email):
    conn = connect_db()
    row = conn.execute(
        "SELECT t.* FROM teams t JOIN team_members tm ON t.id = tm.team_id "
        "WHERE tm.user_email = %s AND tm.status = 'active'", (user_email,)).fetchone()
    conn.close()
    return dict(row) if row else None

def get_team_member_emails(team_id):
    conn = connect_db()
    rows = conn.execute(
        "SELECT user_email FROM team_members WHERE team_id = %s AND status = 'active'",
        (team_id,)).fetchall()
    conn.close()
    return [r['user_email'] for r in rows]

def is_team_admin(user_email, team_id):
    conn = connect_db()
    row = conn.execute(
        "SELECT role FROM team_members WHERE team_id = %s AND user_email = %s AND status = 'active'",
        (team_id, user_email)).fetchone()
    conn.close()
    return bool(row and row['role'] == 'admin')

def get_canonical_components_team(user_id):
    team = get_user_team(user_id)
    conn = connect_db()
    if team:
        emails = get_team_member_emails(team['id']) or [user_id]
        placeholders = ','.join(['%s'] * len(emails))
        rows = conn.execute(
            f"SELECT * FROM generated_components WHERE user_id IN ({placeholders}) AND status='canonical' ORDER BY name ASC",
            emails).fetchall()
    else:
        rows = conn.execute(
            "SELECT * FROM generated_components WHERE user_id = %s AND status='canonical' ORDER BY name ASC",
            (user_id,)).fetchall()
    conn.close()
    return [dict(r) for r in rows]

def get_conventions_team(user_id, limit=100):
    team = get_user_team(user_id)
    conn = connect_db()
    if team:
        emails = get_team_member_emails(team['id']) or [user_id]
        placeholders = ','.join(['%s'] * len(emails))
        rows = conn.execute(
            f"SELECT * FROM conventions WHERE user_id IN ({placeholders}) AND feedback_type!='dismissed' ORDER BY id DESC LIMIT %s",
            emails + [limit]).fetchall()
    else:
        rows = conn.execute(
            "SELECT * FROM conventions WHERE user_id = %s AND feedback_type!='dismissed' ORDER BY id DESC LIMIT %s",
            (user_id, limit)).fetchall()
    conn.close()
    return [dict(r) for r in rows]

def get_ds_files_team(user_id):
    team = get_user_team(user_id)
    conn = connect_db()
    if team:
        emails = get_team_member_emails(team['id']) or [user_id]
        placeholders = ','.join(['%s'] * len(emails))
        rows = conn.execute(
            f"SELECT * FROM ds_files WHERE user_id IN ({placeholders}) ORDER BY id DESC",
            emails).fetchall()
    else:
        rows = conn.execute(
            "SELECT * FROM ds_files WHERE user_id = %s ORDER BY id DESC",
            (user_id,)).fetchall()
    conn.close()
    return [dict(r) for r in rows]

# ============================================================
# GAP RESOLUTIONS
# ============================================================
def save_gap_resolution(user_id, spec_id, gap_title, status, notes=""):
    conn = connect_db()
    conn.execute("DELETE FROM gap_resolutions WHERE user_id=? AND spec_id=? AND gap_title=?",(user_id,spec_id,gap_title))
    conn.execute("INSERT INTO gap_resolutions (user_id,spec_id,gap_title,status,notes,created_at) VALUES (?,?,?,?,?,?)",
        (user_id,spec_id,gap_title,status,notes,datetime.now().strftime("%Y-%m-%d %H:%M")))
    conn.commit(); conn.close()

def get_gap_resolutions(user_id, spec_id):
    conn = connect_db()
    rows = conn.execute("SELECT * FROM gap_resolutions WHERE user_id=? AND spec_id=?",(user_id,spec_id)).fetchall()
    conn.close(); return {r['gap_title']:dict(r) for r in rows}

# ============================================================
# DS FILE HELPERS
# ============================================================
def parse_figma_tokens(raw):
    lines = []
    def walk(obj, path=""):
        if isinstance(obj, dict):
            for k,v in obj.items(): walk(v, f"{path}.{k}" if path else k)
        elif isinstance(obj, list):
            for item in obj: walk(item, path)
        else: lines.append(f"{path}: {obj}")
    walk(raw)
    return "\n".join(lines[:500])

def parse_css_file(text):
    import re
    tokens  = re.findall(r'(--[\w-]+)\s*:\s*([^;]+);', text)
    classes = re.findall(r'\.([\w-]+)\s*\{', text)
    out  = "CSS TOKENS:\n" + "\n".join(f"  {k}: {v.strip()}" for k,v in tokens[:200])
    out += "\n\nCOMPONENT CLASSES:\n" + "\n".join(f"  .{c}" for c in classes[:200])
    return out

def extract_ds_file(filename, file_bytes):
    ext = filename.rsplit('.',1)[-1].lower()
    if ext == 'json':
        try:
            raw = json.loads(file_bytes.decode('utf-8','ignore'))
            return "FIGMA TOKENS:\n" + parse_figma_tokens(raw)
        except: return file_bytes.decode('utf-8','ignore')[:8000]
    elif ext in ['css','scss']: return parse_css_file(file_bytes.decode('utf-8','ignore'))
    elif ext == 'docx':
        try:
            from docx import Document as DocxDoc; import io as _io
            doc = DocxDoc(_io.BytesIO(file_bytes))
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())[:8000]
        except: return "[Could not parse DOCX]"
    else: return file_bytes.decode('utf-8','ignore')[:8000]

def save_ds_file(filename, content_text, user_id="anonymous"):
    conn = connect_db()
    conn.execute("INSERT INTO ds_files (filename,file_type,content,uploaded_at,user_id) VALUES (?,?,?,?,?)",
        (filename, filename.rsplit('.',1)[-1].lower(), content_text, datetime.now().strftime("%Y-%m-%d %H:%M"), user_id))
    conn.commit()
    row = conn.execute("SELECT id FROM ds_files WHERE user_id=? ORDER BY id DESC LIMIT 1", (user_id,)).fetchone()
    conn.close()
    if row: rag.ingest_ds_file(user_id, row['id'], filename, content_text)

def get_ds_files(user_id="anonymous"):
    conn = connect_db()
    rows = conn.execute("SELECT * FROM ds_files WHERE user_id=? ORDER BY id DESC",(user_id,)).fetchall()
    conn.close(); return [dict(r) for r in rows]

def delete_ds_file(file_id, user_id="anonymous"):
    conn = connect_db()
    conn.execute("DELETE FROM ds_files WHERE id=? AND user_id=?", (file_id, user_id))
    conn.commit(); conn.close()

def get_design_system_context(user_id="anonymous"):
    files = get_ds_files_team(user_id)
    if not files: return ""
    ctx = "YOUR DESIGN SYSTEM FILES:\n"
    for f in files: ctx += f"\n--- {f['filename']} ---\n{f['content']}\n"
    return ctx[:12000]

# ============================================================
# RESEARCH HELPERS
# ============================================================
def fetch_url_content(url):
    import urllib.request, re
    headers = {'User-Agent': 'Mozilla/5.0'}
    req = urllib.request.Request(url, headers=headers)
    with urllib.request.urlopen(req, timeout=10) as resp:
        raw = resp.read().decode('utf-8','ignore')
    raw = re.sub(r'<script[^>]*>.*?</script>', '', raw, flags=re.DOTALL)
    raw = re.sub(r'<style[^>]*>.*?</style>',  '', raw, flags=re.DOTALL)
    raw = re.sub(r'<[^>]+>', ' ', raw)
    raw = re.sub(r'\s+', ' ', raw).strip()
    return raw[:8000]

def parse_research_pdf(file_bytes):
    try:
        import fitz
        doc = fitz.open(stream=file_bytes, filetype='pdf')
        return '\n'.join(p.get_text() for p in doc)[:8000]
    except ImportError: return '[PDF parsing requires pymupdf]'
    except Exception as e: return f'[Could not parse PDF: {str(e)}]'

def save_research_item(title, item_type, source, content, user_id="anonymous"):
    conn = connect_db()
    conn.execute('INSERT INTO research_items (title,item_type,source,content,added_at,user_id) VALUES (?,?,?,?,?,?)',
        (title, item_type, source, content, datetime.now().strftime('%Y-%m-%d %H:%M'), user_id))
    conn.commit(); conn.close()

def get_research_items(user_id="anonymous"):
    conn = connect_db()
    rows = conn.execute('SELECT * FROM research_items WHERE user_id=? ORDER BY id DESC',(user_id,)).fetchall()
    conn.close(); return [dict(r) for r in rows]

def delete_research_item(item_id, user_id="anonymous"):
    conn = connect_db()
    conn.execute('DELETE FROM research_items WHERE id=? AND user_id=?', (item_id, user_id))
    conn.commit(); conn.close()

def get_research_context(user_id="anonymous"):
    items = get_research_items(user_id)
    if not items: return ''
    ctx = 'EXTERNAL RESEARCH:\n'
    for item in items:
        ctx += f"\n--- {item['title']} ({item['item_type']}: {item['source']}) ---\n{item['content']}\n"
    return ctx[:10000]

# ============================================================
# CONVENTIONS HELPERS
# ============================================================
def save_convention(user_id, title, description, category, priority, source_ticket, source_screen, feedback_type):
    conn = connect_db()
    conn.execute("INSERT INTO conventions (user_id,title,description,category,priority,source_ticket,source_screen,feedback_type,created_at) VALUES (?,?,?,?,?,?,?,?,?)",
        (user_id, title, description, category, priority, source_ticket, source_screen, feedback_type, datetime.now().strftime("%Y-%m-%d %H:%M")))
    conn.commit()
    row = conn.execute("SELECT id FROM conventions WHERE user_id=? ORDER BY id DESC LIMIT 1", (user_id,)).fetchone()
    conn.close()
    if row:
        rag.ingest_convention(user_id=user_id, convention_id=row['id'], recommendation=description, decision=feedback_type, edited_text=description)

def get_conventions(user_id="anonymous", limit=100):
    conn = connect_db()
    rows = conn.execute("SELECT * FROM conventions WHERE user_id=? AND feedback_type!='dismissed' ORDER BY id DESC LIMIT ?",(user_id, limit)).fetchall()
    conn.close(); return [dict(r) for r in rows]

def delete_convention(conv_id, user_id="anonymous"):
    conn = connect_db()
    conn.execute("DELETE FROM conventions WHERE id=? AND user_id=?", (conv_id, user_id))
    conn.commit(); conn.close()

def get_conventions_context(user_id="anonymous"):
    convs = get_conventions_team(user_id, limit=20)
    if not convs: return ""
    ctx = "DESIGNER'S ESTABLISHED CONVENTIONS:\n"
    for c in convs: ctx += f"  [{c['category']}] {c['title']}: {c['description']}\n"
    ctx += "Prioritise these established patterns.\n"
    return ctx

# ============================================================
# GENERATED COMPONENTS HELPERS
# ============================================================
def get_generated_components(user_id="anonymous"):
    conn = connect_db()
    rows = conn.execute("SELECT * FROM generated_components WHERE user_id=? ORDER BY name ASC",(user_id,)).fetchall()
    conn.close(); return [dict(r) for r in rows]

def get_canonical_components(user_id="anonymous"):
    return get_canonical_components_team(user_id)

def save_generated_component(user_id, name, comp_type, html_code, react_code, source_ticket, source_screen, design_specs=None):
    conn = connect_db()
    specs_json = json.dumps(design_specs) if design_specs else None
    conn.execute("INSERT INTO generated_components (user_id,name,type,html_code,react_code,design_specs,source_ticket,source_screen,status,created_at) VALUES (?,?,?,?,?,?,?,?,'canonical',?)",
        (user_id, name, comp_type, html_code, react_code, specs_json, source_ticket, source_screen, datetime.now().strftime("%Y-%m-%d %H:%M")))
    conn.commit()
    row = conn.execute("SELECT id FROM generated_components WHERE user_id=? ORDER BY id DESC LIMIT 1", (user_id,)).fetchone()
    conn.close()
    if row:
        rag.ingest_component(user_id=user_id, component_id=row['id'], name=name, comp_type=comp_type, design_specs=specs_json or "", status="canonical")
    return {"status": "canonical"}

def resolve_conflict_components(keep_id, discard_id, user_id):
    conn = connect_db()
    conn.execute("UPDATE generated_components SET status='canonical', conflict_with=NULL WHERE id=? AND user_id=?",(keep_id,user_id))
    conn.execute("DELETE FROM generated_components WHERE id=? AND user_id=?",(discard_id,user_id))
    conn.commit(); conn.close()

def delete_generated_component(comp_id, user_id):
    conn = connect_db()
    conn.execute("DELETE FROM generated_components WHERE id=? AND user_id=?",(comp_id,user_id))
    conn.commit(); conn.close()

def get_generated_ds_context(user_id="anonymous"):
    comps = get_canonical_components(user_id)
    if not comps: return ""
    ctx = f"YOUR DESIGN SYSTEM ({len(comps)} confirmed components):\n"
    for c in comps: ctx += f"  - {c['name']} ({c['type']})\n"
    return ctx

# ============================================================
# PRODUCT CONTEXT HELPERS
# ============================================================
def get_product_context_items(user_id="anonymous"):
    conn = connect_db()
    rows = conn.execute("SELECT * FROM product_context WHERE user_id=? ORDER BY id DESC",(user_id,)).fetchall()
    conn.close(); return [dict(r) for r in rows]

def delete_product_context_item(item_id, user_id):
    conn = connect_db()
    conn.execute("DELETE FROM product_context WHERE id=? AND user_id=?",(item_id,user_id))
    conn.commit(); conn.close()

def auto_extract_product_context(analysis, ticket, user_id):
    screen_sum  = analysis.get("screen_summary", "")
    comp_names  = [c.get("name","") for c in analysis.get("components", [])]
    ticket_key  = ticket.get("key", "")
    product_name = analysis.get("product_name", "")
    if not product_name:
        summary_lower = (ticket.get("summary","") + " " + screen_sum).lower()
        product_hints = {"Power BI":["power bi","powerbi","pbi"],"Salesforce":["salesforce","sfdc","crm"],"SAP":["sap","s4 hana"],"Jira":["jira","atlassian"],"Figma":["figma","design tool"]}
        product_name = "Product"
        for pname, hints in product_hints.items():
            if any(h in summary_lower for h in hints): product_name = pname; break
    summary = (ticket.get("summary","") + " " + screen_sum).lower()
    domains = {"Search":["search","filter","query"],"Dashboard":["dashboard","overview","kpi","metric"],"Reports":["report","chart","graph","analytics"],"Navigation":["nav","menu","sidebar","breadcrumb"],"Forms":["form","input","submit","create","edit"],"Authentication":["login","auth","sign in","password"]}
    detected_domain = "General"
    for domain, keywords in domains.items():
        if any(k in summary for k in keywords): detected_domain = domain; break
    conn = connect_db()
    conn.execute("INSERT INTO product_context (user_id,product_name,feature_domain,pattern,insight,source_ticket,source_screen,item_type,content,title,created_at) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
        (user_id, product_name, detected_domain, ", ".join(comp_names[:5]), screen_sum[:200],
         ticket_key, analysis.get("_filename",""), "auto", screen_sum, f"{product_name} — {detected_domain}", datetime.now().strftime("%Y-%m-%d %H:%M")))
    conn.commit()
    row = conn.execute("SELECT id FROM product_context WHERE user_id=? ORDER BY id DESC LIMIT 1", (user_id,)).fetchone()
    conn.close()
    if row: rag.ingest_product_context(user_id=user_id, product_name=product_name, feature_domain=detected_domain, insight=screen_sum[:200], ctx_id=row['id'])

def get_product_context_summary(user_id="anonymous"):
    items = get_product_context_items(user_id)
    if not items: return ""
    from collections import Counter
    domains = Counter(i['feature_domain'] for i in items)
    ctx = "PRODUCT CONTEXT (accumulated from past analyses):\n"
    ctx += f"  Features: {', '.join(f'{d} ({c} screens)' for d,c in domains.most_common(5))}\n"
    for item in items[:5]:
        if item.get('insight'): ctx += f"  - {item['insight'][:100]}\n"
    return ctx[:2000]


# ============================================================
# MEMORY
# ============================================================
def save_analysis_to_memory(analysis):
    ticket_id  = analysis.get("_ticket_id","unknown")
    ticket_sum = analysis.get("_ticket_data",{}).get("summary","")
    screen     = analysis.get("_filename","screen.png")
    compliance = analysis.get("compliance_score",0)
    components = ",".join([c.get("name","") for c in analysis.get("components",[])])
    spec_id    = f"{ticket_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    json_path  = os.path.join(SPECS_DIR, f"{spec_id}.json")
    created_at = datetime.now().strftime("%Y-%m-%d %H:%M")
    try:
        os.makedirs(SPECS_DIR, exist_ok=True)
        with open(json_path,"w") as f: json.dump(analysis,f,indent=2)
    except: pass
    conn = connect_db(); user_id = current_user_id()
    spec_json = json.dumps(analysis)
    if _USE_PG:
        conn.execute("INSERT INTO specs (spec_id,ticket_id,ticket_summary,screen_file,compliance,components,created_at,json_path,user_id,spec_data) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s) ON CONFLICT (spec_id) DO UPDATE SET spec_data=EXCLUDED.spec_data",
            (spec_id,ticket_id,ticket_sum,screen,compliance,components,created_at,json_path,user_id,spec_json))
    else:
        conn.execute("INSERT OR REPLACE INTO specs (spec_id,ticket_id,ticket_summary,screen_file,compliance,components,created_at,json_path,user_id,spec_data) VALUES (?,?,?,?,?,?,?,?,?,?)",
            (spec_id,ticket_id,ticket_sum,screen,compliance,components,created_at,json_path,user_id,spec_json))
    conn.commit(); conn.close()
    return spec_id

def search_specs(query="",search_type="all",date_filter="",user_id=None):
    conn=connect_db()
    base="SELECT * FROM specs"; where=["user_id=?"]; params=[user_id or "anonymous"]
    if query:
        if search_type=="ticket":      where.append("ticket_id LIKE ?");    params.append(f"%{query}%")
        elif search_type=="component": where.append("components LIKE ?");    params.append(f"%{query}%")
        else:
            where.append("(ticket_id LIKE ? OR components LIKE ? OR ticket_summary LIKE ?)")
            params+=[f"%{query}%",f"%{query}%",f"%{query}%"]
    if date_filter: where.append("created_at LIKE ?"); params.append(f"{date_filter}%")
    if where: base+=" WHERE "+" AND ".join(where)
    base+=" ORDER BY id DESC LIMIT 100"
    rows=conn.execute(base,params).fetchall(); conn.close()
    return [dict(r) for r in rows]

def get_recent_specs(limit=5, user_id=None):
    conn=connect_db()
    rows=conn.execute("SELECT * FROM specs WHERE user_id=? ORDER BY id DESC LIMIT ?",(user_id or "anonymous",limit)).fetchall()
    conn.close(); return [dict(r) for r in rows]

def load_spec_by_id(spec_id):
    conn=connect_db()
    row=conn.execute("SELECT json_path FROM specs WHERE spec_id=?",(spec_id,)).fetchone()
    conn.close()
    if not row: return None
    try:
        with open(row["json_path"]) as f: return json.load(f)
    except: return None

# ============================================================
# JIRA
# ============================================================
def fetch_jira_ticket(ticket_id):
    url  = f"https://{JIRA_DOMAIN}/rest/api/3/issue/{ticket_id}"
    auth = b64_lib.b64encode(f"{JIRA_EMAIL}:{JIRA_API_TOKEN}".encode()).decode()
    resp = http_requests.get(url, headers={"Authorization":f"Basic {auth}","Accept":"application/json"}, timeout=10)
    if resp.status_code == 200:
        data=resp.json(); fields=data.get("fields",{})
        def adf(n):
            if not n: return ""
            if isinstance(n,str): return n
            t=n.get("text","") if n.get("type")=="text" else ""
            for c in n.get("content",[]): t+=adf(c)+" "
            return t.strip()
        acc=""
        for k,v in fields.items():
            if "acceptance" in k.lower() and v: acc=adf(v) if isinstance(v,dict) else str(v)
        return {"key":data.get("key",ticket_id),"summary":fields.get("summary",""),"description":adf(fields.get("description",{}))[:3000],"status":fields.get("status",{}).get("name",""),"issue_type":fields.get("issuetype",{}).get("name",""),"priority":fields.get("priority",{}).get("name",""),"assignee":(fields.get("assignee") or {}).get("displayName","Unassigned"),"acceptance_criteria":acc[:1500],"labels":fields.get("labels",[])}
    elif resp.status_code==401: raise Exception("Jira authentication failed.")
    elif resp.status_code==404: raise Exception(f"Ticket '{ticket_id}' not found.")
    else: raise Exception(f"Jira API error {resp.status_code}")

def create_jira_ticket(summary, description, issue_type="Task"):
    url  = f"https://{JIRA_DOMAIN}/rest/api/3/issue"
    auth = b64_lib.b64encode(f"{JIRA_EMAIL}:{JIRA_API_TOKEN}".encode()).decode()
    proj_resp = http_requests.get(f"https://{JIRA_DOMAIN}/rest/api/3/project",
        headers={"Authorization":f"Basic {auth}","Accept":"application/json"}, timeout=10)
    if proj_resp.status_code != 200: raise Exception("Could not fetch Jira projects.")
    projects = proj_resp.json()
    if not projects: raise Exception("No Jira projects found.")
    project_key = projects[0]["key"]
    payload = {"fields":{"project":{"key":project_key},"summary":summary,"description":{"type":"doc","version":1,"content":[{"type":"paragraph","content":[{"type":"text","text":description}]}]},"issuetype":{"name":issue_type}}}
    resp = http_requests.post(url, headers={"Authorization":f"Basic {auth}","Content-Type":"application/json","Accept":"application/json"}, json=payload, timeout=10)
    if resp.status_code in [200,201]: return resp.json()
    else: raise Exception(f"Failed to create ticket: {resp.status_code}")

# ============================================================
# PDF GENERATION
# ============================================================
def generate_pdf(analysis, filename):
    buf=io.BytesIO()
    doc=SimpleDocTemplate(buf,pagesize=letter,leftMargin=inch,rightMargin=inch,topMargin=inch,bottomMargin=inch)
    S=getSampleStyleSheet()
    BLUE=colors.HexColor("#0f62fe");DARK=colors.HexColor("#161616");GREY=colors.HexColor("#6f6f6f")
    LGREY=colors.HexColor("#f4f4f4");GREEN=colors.HexColor("#198038");RED=colors.HexColor("#da1e28")
    GOLD=colors.HexColor("#b28600");ORANGE=colors.HexColor("#ff832b")
    ts=ParagraphStyle('T',parent=S['Normal'],fontSize=22,textColor=DARK,fontName='Helvetica-Bold',spaceAfter=4)
    ss=ParagraphStyle('S',parent=S['Normal'],fontSize=10,textColor=GREY,spaceAfter=20)
    h1=ParagraphStyle('H1',parent=S['Normal'],fontSize=13,textColor=BLUE,fontName='Helvetica-Bold',spaceBefore=18,spaceAfter=6)
    bs=ParagraphStyle('B',parent=S['Normal'],fontSize=10,leading=15,textColor=DARK,spaceAfter=6)
    bl=ParagraphStyle('BL',parent=S['Normal'],fontSize=10,leading=15,textColor=DARK,leftIndent=16,spaceAfter=4)
    st={"Covered":GREEN,"Partial":ORANGE,"Missing":RED}
    sv={"HIGH":RED,"MEDIUM":ORANGE,"LOW":GOLD}
    story=[]
    story.append(Paragraph("Ai4UX — Component Spec + Requirements Audit",ts))
    story.append(Paragraph(f"Generated: {datetime.now().strftime('%d %b %Y, %H:%M')}  |  {filename}  |  {analysis.get('_ticket_id','')}",ss))
    tk=analysis.get("_ticket_data",{})
    if tk:
        story.append(Paragraph("JIRA TICKET",h1))
        tbl=Table([["Key",tk.get("key","")],["Summary",tk.get("summary","")],["Type",tk.get("issue_type","")],["Status",tk.get("status","")],["Priority",tk.get("priority","")]],colWidths=[1.2*inch,5.3*inch])
        tbl.setStyle(TableStyle([('FONTNAME',(0,0),(0,-1),'Helvetica-Bold'),('FONTSIZE',(0,0),(-1,-1),9),('TEXTCOLOR',(0,0),(0,-1),BLUE),('ROWBACKGROUNDS',(0,0),(-1,-1),[colors.white,LGREY]),('GRID',(0,0),(-1,-1),0.5,colors.HexColor("#e0e0e0")),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5),('LEFTPADDING',(0,0),(-1,-1),8)]))
        story.append(tbl); story.append(Spacer(1,8))
    story.append(Paragraph("SCREEN SUMMARY",h1)); story.append(Paragraph(analysis.get("screen_summary",""),bs))
    story.append(Paragraph("JIRA SUMMARY",h1));   story.append(Paragraph(analysis.get("jira_summary",""),bs))
    cs=analysis.get("compliance_score",0)
    story.append(Paragraph("COMPLIANCE",h1)); story.append(Paragraph(f"{cs}%",bs))
    gaps=analysis.get("gaps",[])
    if gaps:
        story.append(Paragraph("GAPS",h1))
        for g in gaps:
            sv_color=sv.get(g.get("severity",""),GREY)
            story.append(Paragraph(f'<font color="#{sv_color.hexval()[2:]}"><b>{g.get("title","")} [{g.get("severity","")}]</b></font>',bs))
            story.append(Paragraph(g.get("description",""),bl))
            story.append(Paragraph(f"Recommendation: {g.get('recommendation','')}",bl))
            story.append(Spacer(1,6))
    story.append(Spacer(1,20))
    story.append(Paragraph("Generated by Ai4UX",ParagraphStyle('F',parent=S['Normal'],fontSize=8,textColor=GREY,alignment=1)))
    doc.build(story); buf.seek(0); return buf


# ============================================================
# DESIGN SYSTEM PROFILES + PROMPTS
# ============================================================
DS_PROFILES = {
    "carbon":   {"name":"Carbon Design System (IBM)","components":"Button, Text Input, Form, Modal, Dropdown, Checkbox, Radio Button, Toggle, Notification, Breadcrumb, Tabs, Tag, Data Table, Card, Search, Pagination, Progress Indicator, Loading, Tooltip, Link, Header, Side Nav, Content Switcher, Accordion","tokens":"@carbon/styles, prefix: cds--. Colors: $blue-60 (#0f62fe), $gray-100 (#161616). Typography: IBM Plex Sans.","code_prefix":"Use Carbon React imports: import { Button } from '@carbon/react'"},
    "material": {"name":"Material Design 3 (Google)","components":"Filled Button, Outlined Button, FAB, Card, Chip, Dialog, Icon Button, List, Menu, Navigation Bar, Navigation Drawer, Progress Indicator, Radio Button, Search, Snackbar, Switch, Tab, Text Field, Top App Bar, Tooltip","tokens":"md-sys-color. Primary: #6750A4. Surface: #FFFBFE.","code_prefix":"Use MUI: import { Button } from '@mui/material'"},
    "ant":      {"name":"Ant Design (Alibaba)","components":"Button, Input, Select, Checkbox, Radio, Switch, DatePicker, Form, Table, List, Card, Collapse, Tabs, Tag, Badge, Alert, Modal, Drawer, Dropdown, Menu, Breadcrumb, Pagination, Steps, Spin, Progress, Avatar, Tooltip","tokens":"Primary: #1677ff. Border radius: 6px.","code_prefix":"Use Ant Design: import { Button } from 'antd'"},
    "shadcn":   {"name":"Shadcn/UI","components":"Accordion, Alert, Alert Dialog, Avatar, Badge, Button, Calendar, Card, Checkbox, Command, Dialog, Drawer, Dropdown Menu, Form, Input, Label, Navigation Menu, Pagination, Popover, Progress, Radio Group, Select, Separator, Sheet, Skeleton, Slider, Switch, Table, Tabs, Textarea, Toggle, Tooltip","tokens":"CSS variables: --background, --foreground, --primary, --secondary, --muted. Uses Tailwind CSS + Radix UI.","code_prefix":"Use shadcn/ui: import { Button } from '@/components/ui/button'"},
    "custom":   {"name":"Custom Design System","components":"Use components from uploaded design system files","tokens":"Reference uploaded token files","code_prefix":"Use component names from your uploaded design system files"},
}

def get_screen_only_prompt(ds_id="carbon"):
    ds = DS_PROFILES.get(ds_id, DS_PROFILES["carbon"])
    return f"""You are an expert UX analyst and frontend developer specialising in the {ds["name"]}.
Analyse this UI screenshot as a standalone screen. Return ONLY valid JSON.
{{
  "analysis_mode": "screen_only",
  "screen_summary": "One sentence describing what this screen does",
  "product_name": "Infer from the screen content",
  "compliance_score": 82,
  "design_system_compliance": {{ "score": "High / Medium / Low", "summary": "One sentence" }},
  "components": [{{
    "name": "ComponentName", "type": "Input/Button/Container/Navigation/Card/Tag/Placeholder",
    "instances": 1, "complexity": "Simple/Medium/Complex",
    "carbon_equivalent": "Exact {ds["name"]} component name", "notes": "Usage or accessibility note",
    "html_code": "Concise HTML snippet max 15 lines",
    "react_code": "Concise JSX snippet max 15 lines. {ds["code_prefix"]}",
    "design_specs": {{
      "primary_color": "#hex", "background": "#hex", "border": "1px solid #e5e7eb",
      "border_radius": "4px", "typography": "font-family, size, weight",
      "spacing": "padding values", "shadow": "box-shadow or none", "states": "hover/focus descriptions"
    }}
  }}],
  "gaps": [{{"title":"Gap title","severity":"HIGH/MEDIUM/LOW","description":"...","recommendation":"...","effort":"1-2 days","jira_reference":"N/A"}}],
  "recommendations": [{{"category":"UX Enhancement/Accessibility/Performance/Design System","priority":"HIGH PRIORITY/MEDIUM/LOW","estimate":"1 day","title":"...","description":"...","benefit":"...","dependencies":"null"}}],
  "requirements_audit": [],
  "citations": [{{"criterion":"WCAG 1.4.3","registry":"WCAG 2.2 AA","title":"Contrast (Minimum)","status":"fail","detail":"...","url":"...","fix":"..."}}]
}}
Design system: {ds["name"]}
Components: {ds["components"]}
Tokens: {ds["tokens"]}
design_specs is MANDATORY for every component."""

def get_system_prompt(ds_id="carbon"):
    ds = DS_PROFILES.get(ds_id, DS_PROFILES["carbon"])
    return f"""You are an expert UX analyst and frontend developer specialising in the {ds["name"]}.
Analyse a UI screenshot alongside a Jira ticket. Respond with ONLY valid JSON.
{{
  "screen_summary": "One sentence describing what this screen does",
  "jira_summary": "One sentence summarising what the Jira ticket requires",
  "compliance_score": 87,
  "product_name": "Infer from ticket + screen. Never leave blank.",
  "design_system_compliance": {{ "score": "High/Medium/Low", "summary": "One sentence" }},
  "requirements_audit": [{{"requirement":"...","status":"Covered/Partial/Missing","location":"..."}}],
  "components": [{{"name":"...","type":"Input/Button/Container/Navigation/Card/Tag/Placeholder","instances":1,"complexity":"Simple/Medium/Complex","carbon_equivalent":"...","notes":"...","html_code":"max 15 lines","react_code":"max 15 lines","design_specs":{{"primary_color":"#hex","background":"#hex","border":"...","border_radius":"...","typography":"...","spacing":"...","shadow":"...","states":"..."}}}}],
  "gaps": [{{"title":"...","severity":"HIGH/MEDIUM/LOW","description":"...","recommendation":"...","effort":"...","jira_reference":"..."}}],
  "recommendations": [{{"category":"UX Enhancement/Accessibility/Performance/Feature/Design System","priority":"HIGH PRIORITY/MEDIUM/LOW","estimate":"...","title":"...","description":"...","benefit":"...","dependencies":"..."}}],
  "citations": [{{"criterion":"WCAG 1.4.3","registry":"WCAG 2.2 AA","title":"...","status":"fail","detail":"...","url":"...","fix":"..."}}]
}}
Design system: {ds["name"]}
Components: {ds["components"]}
Tokens: {ds["tokens"]}
Code: {ds["code_prefix"]}
design_specs is MANDATORY for every component."""


# ============================================================
# FLASK ROUTES
# ============================================================

# ── Emergency DB init ─────────────────────────────────────────
@app.route("/init-db-now")
def init_db_now():
    results = []
    tables = [
        ("specs", "CREATE TABLE IF NOT EXISTS specs (id SERIAL PRIMARY KEY, spec_id TEXT UNIQUE, ticket_id TEXT, ticket_summary TEXT, screen_file TEXT, compliance INTEGER, components TEXT, created_at TEXT, json_path TEXT, user_id TEXT DEFAULT 'anonymous', spec_data TEXT)"),
        ("ds_files", "CREATE TABLE IF NOT EXISTS ds_files (id SERIAL PRIMARY KEY, filename TEXT, file_type TEXT, content TEXT, uploaded_at TEXT, user_id TEXT DEFAULT 'anonymous')"),
        ("research_items", "CREATE TABLE IF NOT EXISTS research_items (id SERIAL PRIMARY KEY, title TEXT, item_type TEXT, source TEXT, content TEXT, added_at TEXT, user_id TEXT DEFAULT 'anonymous')"),
        ("conventions", "CREATE TABLE IF NOT EXISTS conventions (id SERIAL PRIMARY KEY, user_id TEXT DEFAULT 'anonymous', title TEXT, description TEXT, category TEXT, priority TEXT, source_ticket TEXT, source_screen TEXT, feedback_type TEXT, created_at TEXT)"),
        ("generated_components", "CREATE TABLE IF NOT EXISTS generated_components (id SERIAL PRIMARY KEY, user_id TEXT DEFAULT 'anonymous', name TEXT, type TEXT, html_code TEXT, react_code TEXT, design_specs TEXT, source_ticket TEXT, source_screen TEXT, status TEXT DEFAULT 'canonical', conflict_with INTEGER, created_at TEXT)"),
        ("product_context", "CREATE TABLE IF NOT EXISTS product_context (id SERIAL PRIMARY KEY, user_id TEXT DEFAULT 'anonymous', product_name TEXT DEFAULT 'Unknown', feature_domain TEXT, pattern TEXT, insight TEXT, source_ticket TEXT, source_screen TEXT, item_type TEXT DEFAULT 'auto', content TEXT, title TEXT, created_at TEXT)"),
        ("guideline_activation", "CREATE TABLE IF NOT EXISTS guideline_activation (id SERIAL PRIMARY KEY, user_id TEXT DEFAULT 'anonymous', registry_id TEXT, is_active INTEGER DEFAULT 1, scope TEXT DEFAULT 'global', product_name TEXT DEFAULT '', created_at TEXT)"),
        ("custom_guidelines", "CREATE TABLE IF NOT EXISTS custom_guidelines (id SERIAL PRIMARY KEY, user_id TEXT DEFAULT 'anonymous', name TEXT, version TEXT, category TEXT, content TEXT, filename TEXT, is_active INTEGER DEFAULT 1, created_at TEXT)"),
        ("guideline_profiles", "CREATE TABLE IF NOT EXISTS guideline_profiles (id SERIAL PRIMARY KEY, user_id TEXT DEFAULT 'anonymous', product_name TEXT, active_registries TEXT, overrides TEXT, created_at TEXT)"),
        ("conflict_resolutions", "CREATE TABLE IF NOT EXISTS conflict_resolutions (id SERIAL PRIMARY KEY, user_id TEXT DEFAULT 'anonymous', conflict_id TEXT, resolution TEXT, remember INTEGER DEFAULT 0, created_at TEXT)"),
        ("gap_resolutions", "CREATE TABLE IF NOT EXISTS gap_resolutions (id SERIAL PRIMARY KEY, user_id TEXT DEFAULT 'anonymous', spec_id TEXT, gap_title TEXT, status TEXT DEFAULT 'open', notes TEXT, created_at TEXT)"),
        ("teams", "CREATE TABLE IF NOT EXISTS teams (id SERIAL PRIMARY KEY, name TEXT NOT NULL, created_by TEXT NOT NULL, invite_code TEXT UNIQUE NOT NULL, created_at TEXT)"),
        ("team_members", "CREATE TABLE IF NOT EXISTS team_members (id SERIAL PRIMARY KEY, team_id INTEGER NOT NULL, user_email TEXT NOT NULL, role TEXT DEFAULT 'member', status TEXT DEFAULT 'pending', invited_by TEXT, joined_at TEXT, UNIQUE(team_id, user_email))"),
        ("team_invites", "CREATE TABLE IF NOT EXISTS team_invites (id SERIAL PRIMARY KEY, team_id INTEGER NOT NULL, invited_email TEXT NOT NULL, invited_by TEXT NOT NULL, token TEXT UNIQUE NOT NULL, status TEXT DEFAULT 'pending', created_at TEXT)"),
        ("spec_sharing", "CREATE TABLE IF NOT EXISTS spec_sharing (id SERIAL PRIMARY KEY, spec_id TEXT NOT NULL, shared_by TEXT NOT NULL, team_id INTEGER NOT NULL, created_at TEXT, UNIQUE(spec_id, team_id))"),
    ]
    conn = connect_db()
    for name, sql in tables:
        try:
            conn.execute(sql); conn.commit(); results.append(f"OK: {name}")
        except Exception as e:
            results.append(f"ERROR {name}: {str(e)}")
    conn.close()
    return "<br>".join(results) + "<br><br><b>Done.</b>"

# ── RAG ───────────────────────────────────────────────────────
@app.route("/api/rag/stats")
def api_rag_stats():
    user_id = current_user_id()
    stats = rag.stats(user_id)
    if stats.get("available") and stats.get("collections", {}).get("guidelines", 0) == 0:
        try: rag.ingest_builtin_guidelines(BUILTIN_GUIDELINES)
        except: pass
    return jsonify(rag.stats(user_id))

@app.route("/api/rag/reindex", methods=["POST"])
@login_required
def api_rag_reindex():
    _rag_warmup_async(current_user_id())
    return jsonify({"status": "reindex_started"})

# ── Auth ──────────────────────────────────────────────────────
LOGIN_PAGE = """<!DOCTYPE html><html><head><meta charset="UTF-8"><title>Ai4UX</title>
<link href="https://fonts.googleapis.com/css2?family=IBM+Plex+Mono:wght@400;600;700&display=swap" rel="stylesheet">
<style>*{font-family:'IBM Plex Mono',monospace;box-sizing:border-box;margin:0;padding:0;}
body{min-height:100vh;background:#f4f4f4;display:flex;align-items:center;justify-content:center;}
.card{background:white;border-radius:8px;box-shadow:0 4px 24px rgba(0,0,0,0.08);padding:48px 40px;width:400px;text-align:center;}
.logo{font-size:32px;font-weight:700;margin-bottom:8px;}
.tagline{font-size:12px;color:#6b7280;margin-bottom:40px;letter-spacing:0.05em;}
hr{border:none;border-top:1px solid #e5e7eb;margin:28px 0;}
.gbtn{display:flex;align-items:center;justify-content:center;gap:12px;width:100%;padding:14px 24px;background:white;border:1.5px solid #e5e7eb;border-radius:6px;font-size:14px;font-weight:500;cursor:pointer;text-decoration:none;color:#161616;transition:border-color 0.2s;}
.gbtn:hover{border-color:#0f62fe;box-shadow:0 0 0 3px rgba(15,98,254,0.1);}
.footer{margin-top:32px;font-size:11px;color:#9ca3af;line-height:1.6;}</style></head>
<body><div class="card">
<div class="logo"><span style="color:#3A6FF7;">Ae</span><span style="color:#5DD6FF;">the</span><span style="color:#3A6FF7;">ris</span></div>
<div class="tagline">AI-POWERED UX PIPELINE</div><hr>
<a href="/auth/google" class="gbtn">
<svg width="20" height="20" viewBox="0 0 24 24"><path d="M22.56 12.25c0-.78-.07-1.53-.2-2.25H12v4.26h5.92c-.26 1.37-1.04 2.53-2.21 3.31v2.77h3.57c2.08-1.92 3.28-4.74 3.28-8.09z" fill="#4285F4"/><path d="M12 23c2.97 0 5.46-.98 7.28-2.66l-3.57-2.77c-.98.66-2.23 1.06-3.71 1.06-2.86 0-5.29-1.93-6.16-4.53H2.18v2.84C3.99 20.53 7.7 23 12 23z" fill="#34A853"/><path d="M5.84 14.09c-.22-.66-.35-1.36-.35-2.09s.13-1.43.35-2.09V7.07H2.18C1.43 8.55 1 10.22 1 12s.43 3.45 1.18 4.93l2.85-2.22.81-.62z" fill="#FBBC05"/><path d="M12 5.38c1.62 0 3.06.56 4.21 1.64l3.15-3.15C17.45 2.09 14.97 1 12 1 7.7 1 3.99 3.47 2.18 7.07l3.66 2.84c.87-2.6 3.3-4.53 6.16-4.53z" fill="#EA4335"/></svg>
Sign in with Google</a>
<div class="footer">Only authorised team members can access Aetheris.</div>
</div></body></html>"""

@app.route("/login")
def login():
    if 'user' in session: return redirect('/')
    return render_template_string(LOGIN_PAGE)

@app.route("/auth/google")
def auth_google():
    state = _secrets_mod.token_urlsafe(16)
    session['oauth_state'] = state
    params = {'client_id':GOOGLE_CLIENT_ID,'redirect_uri':GOOGLE_REDIRECT_URI,'response_type':'code','scope':'openid email profile','state':state,'access_type':'online'}
    return redirect(GOOGLE_AUTH_URL + '?' + urllib.parse.urlencode(params))

@app.route("/auth/callback")
def auth_callback():
    error = request.args.get('error')
    if error: return f"<p>Login failed: {error}. <a href='/login'>Try again</a></p>"
    code  = request.args.get('code')
    state = request.args.get('state')
    if state != session.get('oauth_state'): return "<p>Invalid state. <a href='/login'>Try again</a></p>"
    token_resp = http_requests.post(GOOGLE_TOKEN_URL, data={'code':code,'client_id':GOOGLE_CLIENT_ID,'client_secret':GOOGLE_CLIENT_SECRET,'redirect_uri':GOOGLE_REDIRECT_URI,'grant_type':'authorization_code'})
    token_data = token_resp.json()
    access_token = token_data.get('access_token')
    if not access_token: return f"<p>Token error. <a href='/login'>Try again</a></p>"
    user_resp = http_requests.get(GOOGLE_USERINFO_URL, headers={'Authorization': f'Bearer {access_token}'})
    user_info = user_resp.json()
    session['user'] = {'email':user_info.get('email'),'name':user_info.get('name'),'picture':user_info.get('picture')}
    session.pop('oauth_state', None)
    _rag_warmup_async(session['user']['email'])
    return redirect('/')

@app.route("/logout")
def logout():
    session.clear(); return redirect('/login')

@app.route("/me")
def me(): return jsonify(session.get('user', {}))

@app.route("/")
@login_required
def index():
    if USE_REACT: return send_from_directory(DIST_DIR, "index.html")
    return "React build not found."

@app.route("/assets/<path:filename>")
def serve_assets(filename):
    return send_from_directory(os.path.join(DIST_DIR, "assets"), filename)

# ── Core routes ───────────────────────────────────────────────
@app.route("/fetch-ticket", methods=["POST"])
def fetch_ticket_route():
    data=request.get_json()
    try: return jsonify(fetch_jira_ticket(data["ticket_id"]))
    except Exception as e: return jsonify({"error":str(e)}),400

@app.route("/analyse", methods=["POST"])
def analyse():
    global last_analysis
    import traceback
    data=request.get_json()
    image_b64=data.get("image"); media_type=data.get("media_type","image/png")
    filename=data.get("filename","screen.png"); ticket=data.get("ticket") or {}
    if not image_b64: return jsonify({"error":"No image provided"}),400
    uid           = current_user_id()
    ds_ctx        = get_design_system_context(uid)
    research_ctx  = get_research_context(uid)
    conv_ctx      = get_conventions_context(uid)
    ds_gen_ctx    = get_generated_ds_context(uid)
    prod_ctx      = get_product_context_summary(uid)
    product_name  = ticket.get("key","").split("-")[0] if ticket.get("key") else None
    guidelines_ctx= get_guidelines_context(uid, product_name)
    screen_description = (ticket.get("summary","") + " " + ticket.get("description","")[:300]).strip()
    rag_ctx = rag.build_analysis_context(components_detected=[], user_id=uid, screen_description=screen_description)
    tc = f"JIRA TICKET:\nKey: {ticket.get('key','')}\nSummary: {ticket.get('summary','')}\nType: {ticket.get('issue_type','')} | Status: {ticket.get('status','')} | Priority: {ticket.get('priority','')}\nDescription: {ticket.get('description','')}\nAcceptance Criteria: {ticket.get('acceptance_criteria','None')}\nLabels: {', '.join(ticket.get('labels',[])) or 'None'}"
    full_ctx = tc
    if rag_ctx:        full_ctx += "\n\n" + rag_ctx
    if ds_ctx:         full_ctx += "\n\n" + ds_ctx
    if research_ctx:   full_ctx += "\n\n" + research_ctx
    if conv_ctx:       full_ctx += "\n\n" + conv_ctx
    if ds_gen_ctx:     full_ctx += "\n\n" + ds_gen_ctx
    if prod_ctx:       full_ctx += "\n\n" + prod_ctx
    if guidelines_ctx: full_ctx += "\n\n" + guidelines_ctx
    try:
        ds_id = data.get("design_system","carbon")
        has_ticket = bool(ticket.get("key","").strip())
        dynamic_prompt = get_system_prompt(ds_id) if has_ticket else get_screen_only_prompt(ds_id)
        msg=client.messages.create(model="claude-sonnet-4-6",max_tokens=16000,system=dynamic_prompt,messages=[{"role":"user","content":[{"type":"image","source":{"type":"base64","media_type":media_type,"data":image_b64}},{"type":"text","text":f"Analyse this UI screen. Return full JSON.\n\n{full_ctx}"}]}])
        raw=msg.content[0].text.strip()
        if raw.startswith("```"): raw=raw.split("```")[1]; raw=raw[4:].strip() if raw.startswith("json") else raw.strip()
        analysis=json.loads(raw)
        analysis["_filename"]=filename; analysis["_ticket_id"]=ticket.get("key",""); analysis["_ticket_data"]=ticket
        if not has_ticket: analysis.setdefault("analysis_mode","screen_only")
        last_analysis=analysis
        save_analysis_to_memory(analysis)
        try: auto_extract_product_context(analysis, ticket, uid)
        except: pass
        existing_names = {c['name'].lower() for c in get_generated_components(uid)}
        analysis['_new_components'] = [c for c in analysis.get('components',[]) if c.get('name','').lower() not in existing_names]
        return jsonify(analysis)
    except json.JSONDecodeError as e:
        return jsonify({"error":f"Could not parse AI response: {str(e)}"}),500
    except Exception as e:
        traceback.print_exc()
        return jsonify({"error":str(e)}),500

@app.route("/create-ticket", methods=["POST"])
def create_ticket_route():
    data=request.get_json()
    try:
        result=create_jira_ticket(data.get("summary",""),data.get("description",""))
        return jsonify({"key":result.get("key","")})
    except Exception as e: return jsonify({"error":str(e)}),400

@app.route("/download-pdf")
def download_pdf():
    global last_analysis
    if not last_analysis: return "No analysis yet.",400
    filename=last_analysis.get("_filename","screen.png")
    return send_file(generate_pdf(last_analysis,filename),mimetype="application/pdf",as_attachment=True,download_name=f"ai4ux_{last_analysis.get('_ticket_id','spec')}.pdf")

@app.route("/add-research-url", methods=["POST"])
def add_research_url():
    data = request.get_json(); url = data.get("url","").strip()
    if not url: return jsonify({"error":"No URL provided"}),400
    try:
        content_text = fetch_url_content(url); title = url.split("//")[-1].split("/")[0]
        save_research_item(title, "URL", url, content_text, current_user_id())
        return jsonify({"ok":True,"title":title,"preview":content_text[:300]})
    except Exception as e: return jsonify({"error":str(e)}),400

@app.route("/upload-research-pdf", methods=["POST"])
def upload_research_pdf():
    if "file" not in request.files: return jsonify({"error":"No file"}),400
    f=request.files["file"]; data=f.read()
    content_text = parse_research_pdf(data)
    save_research_item(f.filename, "PDF", f.filename, content_text, current_user_id())
    return jsonify({"ok":True,"title":f.filename,"preview":content_text[:300]})

@app.route("/research-items")
@login_required
def research_items_route(): return jsonify(get_research_items(current_user_id()))

@app.route("/delete-research/<int:item_id>", methods=["DELETE"])
def delete_research_route(item_id):
    delete_research_item(item_id, current_user_id()); return jsonify({"ok":True})

@app.route("/generated-components")
@login_required
def get_generated_components_route(): return jsonify(get_generated_components(current_user_id()))

@app.route("/save-components", methods=["POST"])
@login_required
def save_components_route():
    data = request.get_json(); uid = current_user_id(); results = []
    for c in data.get("components",[]):
        r = save_generated_component(uid, c.get("name",""), c.get("type",""), c.get("html_code",""), c.get("react_code",""), c.get("source_ticket",""), c.get("source_screen",""), c.get("design_specs"))
        results.append(r)
    return jsonify({"ok":True,"results":results})

@app.route("/resolve-conflict", methods=["POST"])
@login_required
def resolve_conflict_route():
    data=request.get_json(); resolve_conflict_components(data.get("keep_id"),data.get("discard_id"),current_user_id()); return jsonify({"ok":True})

@app.route("/delete-generated/<int:comp_id>", methods=["DELETE"])
@login_required
def delete_generated_route(comp_id):
    uid = current_user_id()
    team = get_user_team(uid)
    if team and not is_team_admin(uid, team['id']):
        comp = next((c for c in get_generated_components(uid) if c['id'] == comp_id), None)
        if comp and comp.get('user_id') != uid:
            return jsonify({"error": "Only team admin can delete shared components"}), 403
    delete_generated_component(comp_id, uid); return jsonify({"ok":True})

@app.route("/product-context")
@login_required
def get_product_context_route(): return jsonify(get_product_context_items(current_user_id()))

@app.route("/product-context/add-url", methods=["POST"])
@login_required
def add_product_context_url():
    data=request.get_json(); url=data.get("url","").strip()
    if not url: return jsonify({"error":"No URL"}),400
    try:
        content_text=fetch_url_content(url); title=url.split("//")[-1].split("/")[0]
        conn=connect_db()
        conn.execute("INSERT INTO product_context (user_id,feature_domain,pattern,insight,source_ticket,source_screen,item_type,content,title,created_at) VALUES (?,?,?,?,?,?,?,?,?,?)",(current_user_id(),"Manual","","",url,"","manual_url",content_text[:8000],title,datetime.now().strftime("%Y-%m-%d %H:%M")))
        conn.commit()
        row = conn.execute("SELECT id FROM product_context WHERE user_id=? ORDER BY id DESC LIMIT 1", (current_user_id(),)).fetchone()
        conn.close()
        if row: rag.ingest_product_context(current_user_id(), title, "Manual", content_text[:200], row['id'])
        return jsonify({"ok":True,"title":title,"preview":content_text[:300]})
    except Exception as e: return jsonify({"error":str(e)}),400

@app.route("/product-context/add-file", methods=["POST"])
@login_required
def add_product_context_file():
    if "file" not in request.files: return jsonify({"error":"No file"}),400
    f=request.files["file"]; data=f.read(); content_text=extract_ds_file(f.filename,data)
    conn=connect_db()
    conn.execute("INSERT INTO product_context (user_id,feature_domain,pattern,insight,source_ticket,source_screen,item_type,content,title,created_at) VALUES (?,?,?,?,?,?,?,?,?,?)",(current_user_id(),"Manual","","",f.filename,"","manual_file",content_text[:8000],f.filename,datetime.now().strftime("%Y-%m-%d %H:%M")))
    conn.commit()
    row = conn.execute("SELECT id FROM product_context WHERE user_id=? ORDER BY id DESC LIMIT 1", (current_user_id(),)).fetchone()
    conn.close()
    if row: rag.ingest_product_context(current_user_id(), f.filename, "Manual", content_text[:200], row['id'])
    return jsonify({"ok":True,"title":f.filename,"preview":content_text[:300]})

@app.route("/product-context/<int:item_id>", methods=["DELETE"])
@login_required
def delete_product_context_route(item_id):
    delete_product_context_item(item_id, current_user_id()); return jsonify({"ok":True})

@app.route("/check-components", methods=["POST"])
@login_required
def check_components_route():
    data = request.get_json(); prompt = data.get("prompt",""); uid = current_user_id()
    if not prompt: return jsonify({"error":"No prompt"}),400
    confirmed = get_canonical_components(uid)
    if not confirmed:
        return jsonify({"can_generate":False,"reason":"empty_ds","message":"Your design system has no confirmed components yet.","missing":[],"matched":[],"confirmed_names":[]})
    check_prompt = f"""A user wants to generate this UI: "{prompt}"
List the UI component types needed. Return ONLY valid JSON:
{{"required_components": ["ComponentName1", "ComponentName2"]}}
Keep names generic (e.g. "Button", "Input", "Table"). Max 8."""
    try:
        msg = client.messages.create(model="claude-sonnet-4-6", max_tokens=300, messages=[{"role":"user","content":check_prompt}])
        raw = msg.content[0].text.strip()
        if raw.startswith("```"): raw=raw.split("```")[1]; raw=raw[4:].strip() if raw.startswith("json") else raw.strip()
        required = json.loads(raw).get("required_components",[])
    except:
        return jsonify({"can_generate":True,"matched":[],"missing":[],"confirmed_names":[c['name'] for c in confirmed]})
    confirmed_names = [c['name'].lower() for c in confirmed]
    def fuzzy_match(req, confirmed_list):
        req_lower = req.lower().replace(" ","").replace("-","").replace("_","")
        for c in confirmed_list:
            c_norm = c.lower().replace(" ","").replace("-","").replace("_","")
            if req_lower in c_norm or c_norm in req_lower: return True
            if set(req.lower().split()) & set(c.lower().split()): return True
        return False
    matched = [r for r in required if fuzzy_match(r, confirmed_names)]
    missing = [r for r in required if not fuzzy_match(r, confirmed_names)]
    return jsonify({"can_generate":len(missing)==0,"reason":"missing_components" if missing else "ok","message":f"Missing {len(missing)} component(s)." if missing else "All components available.","missing":missing,"matched":matched,"confirmed_names":[c['name'] for c in confirmed]})


@app.route("/generate", methods=["POST"])
@login_required
def generate_route():
    data         = request.get_json()
    prompt       = data.get("prompt","")
    uid          = current_user_id()
    use_fallback = data.get("use_fallback", False)
    fallback_ds  = data.get("fallback_ds", "carbon")
    if not prompt: return jsonify({"error":"No prompt"}),400

    ds    = DS_PROFILES.get(fallback_ds, DS_PROFILES["carbon"])
    comps = get_canonical_components(uid)
    prod_ctx = get_product_context_summary(uid)
    rag_ctx  = ""

    if comps:
        ds_list    = "YOUR CONFIRMED DESIGN SYSTEM (use ONLY these components):\n"
        ds_list   += "\n".join(f"  - {c['name']} ({c['type']})" for c in comps)
        code_style = "your custom design system"
        if use_fallback:
            ds_list   += f"\n\nFALLBACK (use for missing components only): {ds['name']} — {ds['tokens']}"
            code_style = f"custom DS + {ds['name']}"
    else:
        if use_fallback:
            ds_list    = f"USE: {ds['name']}\nComponents: {ds['components']}\nTokens: {ds['tokens']}"
            code_style = ds['name']
        else:
            return jsonify({"error":"No confirmed components in your design system. Analyse screens first."}), 400

    if rag_ctx: ds_list += f"\n\nRAG CONTEXT:\n{rag_ctx}"
    if prod_ctx: ds_list += f"\n\n{prod_ctx}"

    structure_prompt = f"""Generate a UI component/screen plan. Return ONLY valid JSON, no markdown.
REQUEST: {prompt}
DESIGN SYSTEM: {ds["name"]}
COMPONENTS AVAILABLE:
{ds_list}
Return ONLY this JSON:
{{
  "title": "short descriptive title",
  "description": "2 sentence description",
  "wireframe": {{
    "title": "Component or Screen Name",
    "sections": [
      {{
        "label": "Section name",
        "layout": "row or column",
        "components": [
          {{"component": "ExactComponentName", "label": "visible text", "width": "full|half|third|auto", "variant": "primary|secondary|default"}}
        ]
      }}
    ]
  }}
}}"""

    figma_prompt = f"""Generate a Figma Plugin API JSON specification for this UI component.
Return ONLY valid JSON — no explanation, no markdown fences, no backticks.
REQUEST: {prompt}
Return exactly this structure:
{{
  "name": "ComponentName",
  "type": "FRAME",
  "width": 400,
  "height": 200,
  "layoutMode": "VERTICAL",
  "primaryAxisSizingMode": "AUTO",
  "counterAxisSizingMode": "FIXED",
  "paddingTop": 24,
  "paddingRight": 24,
  "paddingBottom": 24,
  "paddingLeft": 24,
  "itemSpacing": 16,
  "cornerRadius": 0,
  "fills": [{{"type":"SOLID","r":1.0,"g":1.0,"b":1.0,"a":1.0}}],
  "strokes": [{{"type":"SOLID","r":0.878,"g":0.878,"b":0.878,"a":1.0}}],
  "strokeWeight": 1,
  "description": "Component purpose and accessibility notes",
  "children": [
    {{"type":"TEXT","name":"Heading","characters":"Title","fontSize":14,"fills":[{{"type":"SOLID","r":0.086,"g":0.086,"b":0.086,"a":1.0}}]}}
  ]
}}
Rules: r/g/b are 0.0-1.0. Include 2-5 realistic children. No trailing commas."""

    html_prompt = f"""Generate a complete HTML page for this UI. Return ONLY raw HTML starting with <!DOCTYPE html>.
No explanation, no markdown, no backticks.
REQUEST: {prompt}
DESIGN SYSTEM: {ds["name"]}
COMPONENTS: {ds_list}
Requirements:
- Start with <!DOCTYPE html><html lang="en">
- Use Google Fonts: IBM Plex Sans (weights 300,400,500,600)
- CSS custom properties: --cds-blue-60:#0f62fe; --cds-text-primary:#161616; --cds-background:#ffffff; --cds-layer-01:#f4f4f4; --cds-border-subtle:#e0e0e0
- WCAG AA: focus-visible outlines, proper heading hierarchy, ARIA labels
- Semantic HTML: header, main, section, article
- Complete — do not truncate"""

    react_prompt = f"""Generate a complete React TypeScript component. Return ONLY raw JSX/TSX code.
No explanation, no markdown, no backticks. Start directly with imports.
REQUEST: {prompt}
DESIGN SYSTEM: {ds["name"]}
{ds.get("code_prefix","")}
Requirements:
- Import from '@carbon/react' for all UI components
- Use Carbon CSS class names (cds--tile, cds--btn etc)
- Proper heading hierarchy, ARIA attributes
- Export default function ComponentName()
- TypeScript interfaces for all props
- Complete — do not truncate"""

    storybook_prompt = f"""Generate a Storybook story file for this component.
Return ONLY raw TypeScript code, no markdown fences.
COMPONENT: {prompt}
DESIGN SYSTEM: {ds["name"]}
Include: default export with title/component, at least 3 named story exports (Default, Primary, WithError or similar), args for each story.
Use CSF3 format. Start with: import type {{ Meta, StoryObj }} from '@storybook/react';"""

    readme_prompt = f"""Generate a README.md for this UI component.
Return ONLY raw markdown, no fences.
COMPONENT: {prompt}
DESIGN SYSTEM: {ds["name"]}
Include sections: Overview, Props table (Name | Type | Default | Description), Usage example, Accessibility notes, Design tokens used.
Keep it concise — under 60 lines."""

    try:
        import concurrent.futures
        def call_claude(p, tokens=4000):
            return client.messages.create(
                model="claude-sonnet-4-6", max_tokens=tokens,
                messages=[{"role":"user","content":p}]
            ).content[0].text.strip()

        with concurrent.futures.ThreadPoolExecutor(max_workers=6) as ex:
            f1 = ex.submit(call_claude, structure_prompt,  1000)
            f2 = ex.submit(call_claude, html_prompt,       4000)
            f3 = ex.submit(call_claude, react_prompt,      4000)
            f4 = ex.submit(call_claude, figma_prompt,      2000)
            f5 = ex.submit(call_claude, storybook_prompt,  2000)
            f6 = ex.submit(call_claude, readme_prompt,     2000)
            struct_raw    = f1.result(timeout=120)
            html_raw      = f2.result(timeout=120)
            react_raw     = f3.result(timeout=120)
            figma_raw     = f4.result(timeout=120)
            storybook_raw = f5.result(timeout=120)
            readme_raw    = f6.result(timeout=120)

        def strip_fences(s):
            s = s.strip()
            if s.startswith("```"):
                lines = s.split("\n")[1:]
                if lines and lines[-1].strip() == "```": lines = lines[:-1]
                s = "\n".join(lines).strip()
            return s

        if struct_raw.startswith("```"):
            struct_raw = struct_raw.split("```")[1]
            struct_raw = struct_raw[4:].strip() if struct_raw.startswith("json") else struct_raw.strip()
        structure = json.loads(struct_raw)

        html_raw      = strip_fences(html_raw)
        react_raw     = strip_fences(react_raw)
        storybook_raw = strip_fences(storybook_raw)
        readme_raw    = strip_fences(readme_raw)

        if len(html_raw) < 100:
            html_raw = f"""<!DOCTYPE html><html lang="en"><head>
<meta charset="UTF-8"><meta name="viewport" content="width=device-width,initial-scale=1.0">
<link href="https://fonts.googleapis.com/css2?family=IBM+Plex+Sans:wght@300;400;500;600&display=swap" rel="stylesheet">
<style>:root{{--cds-blue-60:#0f62fe;--cds-text-primary:#161616;--cds-background:#ffffff;--cds-layer-01:#f4f4f4;--cds-border-subtle:#e0e0e0;}}
*{{box-sizing:border-box;margin:0;padding:0;font-family:'IBM Plex Sans',sans-serif;}}
body{{background:var(--cds-layer-01);padding:32px;}}
.card{{background:var(--cds-background);border:1px solid var(--cds-border-subtle);padding:24px;margin-bottom:16px;}}
h1{{font-size:2rem;font-weight:300;color:var(--cds-text-primary);margin-bottom:8px;}}
label{{display:block;font-size:0.875rem;font-weight:500;margin-bottom:4px;}}
input{{width:100%;padding:8px 12px;border:1px solid var(--cds-border-subtle);border-radius:0;font-size:14px;margin-bottom:16px;}}
input:focus{{outline:2px solid var(--cds-blue-60);}}
button{{background:var(--cds-blue-60);color:white;border:none;padding:12px 24px;font-size:14px;font-weight:600;cursor:pointer;}}
</style></head><body><main id="main"><div class="card"><h1>{prompt}</h1>
<label for="f1">Field 1</label><input id="f1" type="text" aria-required="true">
<button type="submit">Submit</button></div></main></body></html>"""

        figma_raw_clean = strip_fences(figma_raw)
        try:
            brace_end = figma_raw_clean.rfind("}")
            if brace_end > 0: figma_raw_clean = figma_raw_clean[:brace_end+1]
            figma_spec = json.loads(figma_raw_clean)
        except:
            figma_spec = {}

        confirmed_names = [c['name'].lower() for c in get_canonical_components(uid)]
        ds_compliance = {"score": 0, "matched": [], "unmatched": [], "total": 0}
        if confirmed_names and react_raw:
            import re
            used_comps = list(set(re.findall(r'<([A-Z][A-Za-z0-9]+)', react_raw)))
            matched   = [c for c in used_comps if any(c.lower() in n or n in c.lower() for n in confirmed_names)]
            skip_tags = {'React','Fragment','div','span','main','header','nav','section','article','h1','h2','h3','p','button','input','form','label','ul','li','a','img','svg','path','style'}
            unmatched = [c for c in used_comps if c not in matched and c not in skip_tags]
            total     = len(matched) + len(unmatched)
            score     = round((len(matched)/total)*100) if total>0 else 100
            ds_compliance = {"score": score, "matched": matched, "unmatched": unmatched, "total": total}

        return jsonify({
            "title":          structure.get("title","Generated Component"),
            "description":    structure.get("description",""),
            "wireframe":      structure.get("wireframe",{}),
            "html_code":      html_raw,
            "react_code":     react_raw,
            "figma_spec":     figma_spec,
            "storybook_code": storybook_raw,
            "readme_content": readme_raw,
            "ds_compliance":  ds_compliance,
        })
    except json.JSONDecodeError as e:
        return jsonify({"error":"Could not parse structure. Try a simpler prompt."}),500
    except Exception as e:
        print(f"GENERATE ERROR: {str(e)}")
        return jsonify({"error":str(e)}),500


# ── Gap routes ────────────────────────────────────────────────
@app.route("/gap/resolve", methods=["POST"])
@login_required
def resolve_gap():
    data = request.get_json(); uid = current_user_id()
    save_gap_resolution(uid, data.get("spec_id",""), data.get("gap_title",""), data.get("status","open"), data.get("notes",""))
    return jsonify({"ok":True})

@app.route("/gap/resolutions/<spec_id>")
@login_required
def get_gap_resolutions_route(spec_id):
    return jsonify(get_gap_resolutions(current_user_id(), spec_id))

# ── Guidelines routes ─────────────────────────────────────────
@app.route("/guidelines/registries")
@login_required
def get_registries():
    uid = current_user_id(); active_ids = set(get_active_registries(uid))
    result = []
    for rid, reg in BUILTIN_GUIDELINES.items():
        result.append({"id":rid,"name":reg["name"],"version":reg["version"],"category":reg["category"],"description":reg["description"],"url":reg["url"],"criteria_count":len(reg["criteria"]),"is_active":rid in active_ids,"is_builtin":True})
    conn = connect_db()
    custom = conn.execute("SELECT * FROM custom_guidelines WHERE user_id=?",(uid,)).fetchall()
    conn.close()
    for c in custom: result.append({**dict(c), "is_builtin": False})
    return jsonify(result)

@app.route("/guidelines/toggle", methods=["POST"])
@login_required
def toggle_registry():
    data = request.get_json(); uid = current_user_id()
    set_registry_active(uid, data.get("registry_id"), data.get("is_active",True), data.get("scope","global"), data.get("product_name",""))
    return jsonify({"ok":True})

@app.route("/guidelines/conflicts")
@login_required
def get_conflicts():
    return jsonify(get_active_conflicts(current_user_id()))

@app.route("/guidelines/resolve-conflict", methods=["POST"])
@login_required
def resolve_conflict_guideline():
    data = request.get_json(); uid = current_user_id()
    save_conflict_resolution(uid, data.get("conflict_id"), data.get("resolution"), data.get("remember",False))
    return jsonify({"ok":True})

@app.route("/guidelines/upload", methods=["POST"])
@login_required
def upload_guideline():
    if "file" not in request.files: return jsonify({"error":"No file"}),400
    f = request.files["file"]; data = f.read()
    content_text = extract_ds_file(f.filename, data)
    name = request.form.get("name", f.filename.rsplit(".",1)[0])
    category = request.form.get("category","custom")
    conn = connect_db()
    conn.execute("INSERT INTO custom_guidelines (user_id,name,version,category,content,filename,is_active,created_at) VALUES (?,?,?,?,?,?,1,?)",
        (current_user_id(), name, "1.0", category, content_text[:10000], f.filename, datetime.now().strftime("%Y-%m-%d %H:%M")))
    conn.commit(); conn.close()
    return jsonify({"ok":True,"name":name,"preview":content_text[:200]})

@app.route("/guidelines/custom/<int:item_id>", methods=["DELETE"])
@login_required
def delete_custom_guideline(item_id):
    conn = connect_db()
    conn.execute("DELETE FROM custom_guidelines WHERE id=? AND user_id=?",(item_id,current_user_id()))
    conn.commit(); conn.close()
    return jsonify({"ok":True})

@app.route("/guidelines/profiles")
@login_required
def get_profiles():
    conn = connect_db()
    rows = conn.execute("SELECT * FROM guideline_profiles WHERE user_id=?",(current_user_id(),)).fetchall()
    conn.close(); return jsonify([dict(r) for r in rows])

@app.route("/guidelines/profiles/save", methods=["POST"])
@login_required
def save_profile():
    data = request.get_json(); uid = current_user_id()
    product = data.get("product_name",""); registries = json.dumps(data.get("active_registries",[])); overrides = json.dumps(data.get("overrides",{}))
    conn = connect_db()
    conn.execute("DELETE FROM guideline_profiles WHERE user_id=? AND product_name=?",(uid,product))
    conn.execute("INSERT INTO guideline_profiles (user_id,product_name,active_registries,overrides,created_at) VALUES (?,?,?,?,?)",
        (uid, product, registries, overrides, datetime.now().strftime("%Y-%m-%d %H:%M")))
    conn.commit(); conn.close()
    return jsonify({"ok":True})

@app.route("/guidelines/criteria/<registry_id>")
@login_required
def get_criteria(registry_id):
    if registry_id not in BUILTIN_GUIDELINES: return jsonify({"error":"Not found"}),404
    reg = BUILTIN_GUIDELINES[registry_id]
    return jsonify({"id":registry_id,"name":reg["name"],"criteria":[{"id":k,**v} for k,v in reg["criteria"].items()]})

# ── Compliance certificate ────────────────────────────────────
@app.route("/compliance-certificate", methods=["POST"])
@login_required
def compliance_certificate():
    data = request.get_json()
    title = data.get("title","Generated Component"); description = data.get("description","")
    ds_compliance = data.get("ds_compliance",{}); guidelines = data.get("guidelines",[]); prompt_used = data.get("prompt","")
    uid = current_user_id()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, leftMargin=inch, rightMargin=inch, topMargin=inch, bottomMargin=inch)
    S = getSampleStyleSheet()
    BLUE = colors.HexColor("#0f62fe"); DARK = colors.HexColor("#161616"); GREY = colors.HexColor("#6f6f6f")
    LGREY = colors.HexColor("#f4f4f4"); GREEN = colors.HexColor("#198038"); RED = colors.HexColor("#da1e28")
    ts = ParagraphStyle('T',parent=S['Normal'],fontSize=22,textColor=DARK,fontName='Helvetica-Bold',spaceAfter=4)
    ss = ParagraphStyle('S',parent=S['Normal'],fontSize=10,textColor=GREY,spaceAfter=16)
    h1 = ParagraphStyle('H1',parent=S['Normal'],fontSize=13,textColor=BLUE,fontName='Helvetica-Bold',spaceBefore=16,spaceAfter=6)
    bs = ParagraphStyle('B',parent=S['Normal'],fontSize=10,leading=15,textColor=DARK,spaceAfter=4)
    bl = ParagraphStyle('BL',parent=S['Normal'],fontSize=10,leading=15,textColor=DARK,leftIndent=16,spaceAfter=3)
    story = []
    story.append(Paragraph("Ai4UX Compliance Certificate", ts))
    story.append(Paragraph(f"Generated: {datetime.now().strftime('%d %b %Y, %H:%M')}  ·  {uid}", ss))
    story.append(Paragraph("COMPONENT", h1))
    story.append(Paragraph(f"<b>{title}</b>", bs))
    story.append(Paragraph(description, bs))
    if prompt_used:
        story.append(Paragraph(f"Prompt: {prompt_used[:200]}", ParagraphStyle('sm',parent=S['Normal'],fontSize=9,textColor=GREY,spaceAfter=8)))
    score = ds_compliance.get("score",0)
    score_color = GREEN if score>=80 else RED
    story.append(Paragraph("DESIGN SYSTEM COMPLIANCE", h1))
    story.append(Paragraph(f'<font color="#{score_color.hexval()[2:]}"><b>{score}%</b></font> — {len(ds_compliance.get("matched",[]))} of {ds_compliance.get("total",0)} components matched', bs))
    story.append(Paragraph("ACTIVE GUIDELINES", h1))
    for g in guidelines:
        story.append(Paragraph(f'✓ {g.get("name","")} — {g.get("criteria_count",0)} criteria applied', bl))
    story.append(Spacer(1,16))
    story.append(Paragraph("Generated by Ai4UX — Design Intelligence",
        ParagraphStyle('F',parent=S['Normal'],fontSize=8,textColor=GREY,alignment=1)))
    doc.build(story); buf.seek(0)
    return send_file(buf, mimetype="application/pdf", as_attachment=True,
        download_name=f"ai4ux_compliance_{title.lower().replace(' ','_')}.pdf")

# ── Feedback / Conventions ────────────────────────────────────
@app.route("/conventions")
@login_required
def get_conventions_route(): return jsonify(get_conventions_team(current_user_id()))

@app.route("/feedback/accept", methods=["POST"])
@login_required
def feedback_accept():
    d=request.get_json(); save_convention(current_user_id(),d.get("title",""),d.get("description",""),d.get("category","General"),d.get("priority","MEDIUM"),d.get("source_ticket",""),d.get("source_screen",""),"accepted"); return jsonify({"ok":True})

@app.route("/feedback/edit", methods=["POST"])
@login_required
def feedback_edit():
    d=request.get_json(); save_convention(current_user_id(),d.get("title",""),d.get("description",""),d.get("category","General"),d.get("priority","MEDIUM"),d.get("source_ticket",""),d.get("source_screen",""),"edited"); return jsonify({"ok":True})

@app.route("/feedback/dismiss", methods=["POST"])
@login_required
def feedback_dismiss():
    d=request.get_json(); save_convention(current_user_id(),d.get("title",""),d.get("description",""),d.get("category","General"),d.get("priority","MEDIUM"),d.get("source_ticket",""),d.get("source_screen",""),"dismissed"); return jsonify({"ok":True})

@app.route("/delete-convention/<int:conv_id>", methods=["DELETE"])
@login_required
def delete_convention_route(conv_id):
    uid = current_user_id()
    team = get_user_team(uid)
    if team and not is_team_admin(uid, team['id']):
        conv = next((c for c in get_conventions(uid) if c['id'] == conv_id), None)
        if conv and conv.get('user_id') != uid:
            return jsonify({"error":"Only team admin can delete shared conventions"}),403
    delete_convention(conv_id, uid); return jsonify({"ok":True})

# ── DS files ──────────────────────────────────────────────────
@app.route("/upload-ds", methods=["POST"])
def upload_ds():
    if 'file' not in request.files: return jsonify({"error":"No file"}),400
    f=request.files['file']; file_bytes=f.read(); content_text=extract_ds_file(f.filename,file_bytes)
    save_ds_file(f.filename,content_text,current_user_id()); return jsonify({"ok":True,"filename":f.filename,"preview":content_text[:300]})

@app.route("/ds-files")
@login_required
def ds_files(): return jsonify(get_ds_files_team(current_user_id()))

@app.route("/delete-ds/<int:file_id>", methods=["DELETE"])
def delete_ds(file_id):
    delete_ds_file(file_id, current_user_id()); return jsonify({"ok":True})

# ── Dashboard / Search / Spec ─────────────────────────────────
@app.route("/dashboard-stats")
@login_required
def dashboard_stats():
    uid=current_user_id(); conn=connect_db()
    specs=[dict(s) for s in conn.execute("SELECT * FROM specs WHERE user_id=? ORDER BY id DESC",(uid,)).fetchall()]
    total=len(specs); avg_comp=round(sum(s['compliance'] or 0 for s in specs)/total) if total else 0
    trend=[{"label":s['ticket_id'] or s['spec_id'][:8],"compliance":s['compliance'] or 0} for s in specs[:10]][::-1]
    from collections import Counter; comp_counter=Counter(); gap_counter=Counter(); total_comps=0; total_gaps=0
    for s in specs:
        comps = [c.strip() for c in (s.get('components','') or '').split(',') if c.strip()]
        total_comps += len(comps)
        for c in comps: comp_counter[c] += 1
    for s in specs[:20]:
        try:
            path=s.get('json_path','')
            if path and os.path.exists(path):
                with open(path) as f: d=json.load(f)
                for g in d.get('gaps',[]): gap_counter[g.get('title','')] += 1; total_gaps+=1
        except: pass
    conn.close()
    return jsonify({"total_analyses":total,"avg_compliance":avg_comp,"total_components":total_comps,"total_gaps":total_gaps,"recent":specs[:5],"compliance_trend":trend,"top_components":[{"name":k,"count":v} for k,v in comp_counter.most_common(12)],"top_gaps":[{"title":k,"count":v} for k,v in gap_counter.most_common(6)]})

@app.route("/recent")
@login_required
def recent(): return jsonify(get_recent_specs(5, current_user_id()))

@app.route("/search")
@login_required
def search():
    return jsonify(search_specs(request.args.get("q",""),request.args.get("type","all"),request.args.get("date",""),current_user_id()))

@app.route("/spec/<spec_id>")
def get_spec(spec_id):
    data=load_spec_by_id(spec_id)
    if not data: return jsonify({"error":"Spec not found"}),404
    return jsonify(data)

@app.route("/get-tokens")
@login_required
def get_tokens():
    uid = current_user_id(); comps = get_canonical_components(uid)
    colors_set = {}; typography_set = {}; spacing_set = {}; border_set = {}
    for c in comps:
        if not c.get('design_specs'): continue
        try:
            specs = json.loads(c['design_specs']) if isinstance(c['design_specs'],str) else c['design_specs']
            if specs.get('primary_color'):  colors_set[specs['primary_color']]  = {'value':specs['primary_color'],'usage':c['name']}
            if specs.get('background'):     colors_set[specs.get('background')] = {'value':specs['background'],'usage':c['name']}
            if specs.get('typography'):     typography_set[c['name']]           = specs['typography']
            if specs.get('spacing'):        spacing_set[c['name']]              = specs['spacing']
            if specs.get('border'):         border_set[c['name']]               = specs['border']
        except: pass
    return jsonify({"colors":list(colors_set.values()),"typography":[{"component":k,"value":v} for k,v in typography_set.items()],"spacing":[{"component":k,"value":v} for k,v in spacing_set.items()],"borders":[{"component":k,"value":v} for k,v in border_set.items()]})

@app.route("/get-page-types")
@login_required
def get_page_types():
    uid = current_user_id(); items = get_product_context_items(uid)
    auto = [i for i in items if i['item_type']=='auto']
    from collections import defaultdict
    by_product = defaultdict(lambda: defaultdict(list))
    for item in auto:
        pname = item.get('product_name','Product') or 'Product'
        domain = item.get('feature_domain','General') or 'General'
        by_product[pname][domain].append({'ticket':item.get('source_ticket',''),'insight':item.get('insight','')[:120],'pattern':item.get('pattern',''),'date':item.get('created_at','')})
    result = []
    for product, domains in by_product.items():
        result.append({'product':product,'domains':{d:items for d,items in domains.items()}})
    return jsonify(result)

# ── Team routes ───────────────────────────────────────────────
@app.route("/api/teams/create", methods=["POST"])
@login_required
def create_team():
    if current_user_id() != ADMIN_EMAIL:
        return jsonify({"error": "Only the app admin can create teams"}), 403
    data = request.get_json(); name = data.get("name", "").strip()
    if not name: return jsonify({"error": "Team name required"}), 400
    invite_code = _secrets_mod.token_urlsafe(16)
    conn = connect_db()
    conn.execute("INSERT INTO teams (name, created_by, invite_code, created_at) VALUES (%s, %s, %s, %s)",
        (name, current_user_id(), invite_code, datetime.now().strftime("%Y-%m-%d %H:%M")))
    conn.commit()
    row = conn.execute("SELECT * FROM teams WHERE invite_code = %s", (invite_code,)).fetchone()
    conn.close()
    team_id = dict(row)['id']
    conn2 = connect_db()
    conn2.execute("INSERT INTO team_members (team_id, user_email, role, status, joined_at) VALUES (%s, %s, 'admin', 'active', %s) ON CONFLICT DO NOTHING",
        (team_id, current_user_id(), datetime.now().strftime("%Y-%m-%d %H:%M")))
    conn2.commit(); conn2.close()
    return jsonify({"ok": True, "team": dict(row), "invite_code": invite_code})

@app.route("/api/teams/my")
@login_required
def my_team():
    team = get_user_team(current_user_id())
    if not team: return jsonify({"team": None})
    conn = connect_db()
    members = conn.execute("SELECT user_email, role, status, joined_at FROM team_members WHERE team_id = %s", (team['id'],)).fetchall()
    conn.close()
    team['members'] = [dict(m) for m in members]
    team['is_admin'] = is_team_admin(current_user_id(), team['id'])
    if not team['is_admin']: team.pop('invite_code', None)
    return jsonify({"team": team})

@app.route("/api/teams/invite", methods=["POST"])
@login_required
def invite_to_team():
    uid = current_user_id(); team = get_user_team(uid)
    if not team: return jsonify({"error": "You are not in a team"}), 400
    if not is_team_admin(uid, team['id']): return jsonify({"error": "Only team admins can invite members"}), 403
    data = request.get_json(); email = data.get("email", "").strip().lower()
    if not email: return jsonify({"error": "Email required"}), 400
    token = _secrets_mod.token_urlsafe(32)
    conn = connect_db()
    existing = conn.execute("SELECT * FROM team_members WHERE team_id = %s AND user_email = %s", (team['id'], email)).fetchone()
    if existing: conn.close(); return jsonify({"error": f"{email} is already in this team"}), 400
    conn.execute("INSERT INTO team_invites (team_id, invited_email, invited_by, token, status, created_at) VALUES (%s, %s, %s, %s, 'pending', %s) ON CONFLICT (token) DO NOTHING",
        (team['id'], email, uid, token, datetime.now().strftime("%Y-%m-%d %H:%M")))
    conn.commit(); conn.close()
    invite_url = f"https://web-production-28dbd.up.railway.app/join-team?token={token}"
    return jsonify({"ok": True, "invite_url": invite_url, "invited_email": email, "message": f"Share this link with {email}: {invite_url}"})

@app.route("/join-team")
@login_required
def join_team():
    token = request.args.get("token", "")
    if not token: return "<p>Invalid invite link.</p>"
    uid = current_user_id(); conn = connect_db()
    invite = conn.execute("SELECT * FROM team_invites WHERE token = %s AND status = 'pending'", (token,)).fetchone()
    if not invite: conn.close(); return "<p>Invite link is invalid or already used.</p>"
    invite = dict(invite)
    existing_team = get_user_team(uid)
    if existing_team: conn.close(); return f"<p>You are already in a team ({existing_team['name']}). Leave it first.</p>"
    conn.execute("INSERT INTO team_members (team_id, user_email, role, status, invited_by, joined_at) VALUES (%s, %s, 'member', 'active', %s, %s) ON CONFLICT (team_id, user_email) DO UPDATE SET status='active', joined_at=%s",
        (invite['team_id'], uid, invite['invited_by'], datetime.now().strftime("%Y-%m-%d %H:%M"), datetime.now().strftime("%Y-%m-%d %H:%M")))
    conn.execute("UPDATE team_invites SET status = 'accepted' WHERE token = %s", (token,))
    conn.commit()
    team = conn.execute("SELECT name FROM teams WHERE id = %s", (invite['team_id'],)).fetchone()
    conn.close()
    team_name = dict(team)['name'] if team else "the team"
    return redirect(f"/?joined={team_name}")

@app.route("/api/teams/leave", methods=["POST"])
@login_required
def leave_team():
    uid = current_user_id(); team = get_user_team(uid)
    if not team: return jsonify({"error": "You are not in a team"}), 400
    if is_team_admin(uid, team['id']): return jsonify({"error": "Admins cannot leave. Transfer admin first."}), 400
    conn = connect_db()
    conn.execute("DELETE FROM team_members WHERE team_id = %s AND user_email = %s", (team['id'], uid))
    conn.commit(); conn.close()
    return jsonify({"ok": True})

@app.route("/api/teams/remove-member", methods=["POST"])
@login_required
def remove_team_member():
    uid = current_user_id(); team = get_user_team(uid)
    if not team or not is_team_admin(uid, team['id']): return jsonify({"error": "Admin only"}), 403
    data = request.get_json(); email = data.get("email", "")
    if email == uid: return jsonify({"error": "Cannot remove yourself"}), 400
    conn = connect_db()
    conn.execute("DELETE FROM team_members WHERE team_id = %s AND user_email = %s", (team['id'], email))
    conn.commit(); conn.close()
    return jsonify({"ok": True})

@app.route("/api/specs/share", methods=["POST"])
@login_required
def share_spec():
    uid = current_user_id(); team = get_user_team(uid)
    if not team: return jsonify({"error": "You are not in a team"}), 400
    data = request.get_json(); spec_id = data.get("spec_id", ""); share = data.get("share", True)
    conn = connect_db()
    if share:
        conn.execute("INSERT INTO spec_sharing (spec_id, shared_by, team_id, created_at) VALUES (%s, %s, %s, %s) ON CONFLICT DO NOTHING",
            (spec_id, uid, team['id'], datetime.now().strftime("%Y-%m-%d %H:%M")))
    else:
        conn.execute("DELETE FROM spec_sharing WHERE spec_id = %s AND team_id = %s AND shared_by = %s", (spec_id, team['id'], uid))
    conn.commit(); conn.close()
    return jsonify({"ok": True})

@app.route("/api/specs/team")
@login_required
def team_specs():
    uid = current_user_id(); team = get_user_team(uid)
    if not team: return jsonify([])
    conn = connect_db()
    rows = conn.execute(
        "SELECT s.*, ss.shared_by FROM specs s JOIN spec_sharing ss ON s.spec_id = ss.spec_id WHERE ss.team_id = %s ORDER BY s.id DESC",
        (team['id'],)).fetchall()
    conn.close()
    return jsonify([dict(r) for r in rows])

@app.route("/api/teams/all")
@login_required
def all_teams():
    if current_user_id() != ADMIN_EMAIL: return jsonify({"error": "Admin only"}), 403
    conn = connect_db()
    rows = conn.execute("SELECT * FROM teams ORDER BY id DESC").fetchall()
    conn.close()
    return jsonify([dict(r) for r in rows])

# ── Figma routes ──────────────────────────────────────────────
@app.route("/api/figma/push", methods=["POST"])
@login_required
def figma_push():
    if not FIGMA_TOKEN: return jsonify({"error": "FIGMA_TOKEN not set"}), 500
    if not FIGMA_FILE_KEY: return jsonify({"error": "FIGMA_FILE_KEY not set"}), 500
    data = request.get_json()
    figma_spec = data.get("figma_spec", {}); comp_name = data.get("component_name", "Ai4UX Component")
    if not figma_spec: return jsonify({"error": "No figma_spec provided"}), 400
    headers = {"X-Figma-Token": FIGMA_TOKEN, "Content-Type": "application/json"}
    try:
        file_resp = http_requests.get(f"{FIGMA_API_BASE}/files/{FIGMA_FILE_KEY}?depth=1", headers=headers, timeout=15)
        if file_resp.status_code == 403: return jsonify({"error": "Figma token invalid or no access to file"}), 403
        if file_resp.status_code != 200: return jsonify({"error": f"Figma API error: {file_resp.status_code}"}), 500
        pages = file_resp.json().get("document", {}).get("children", [])
        if not pages: return jsonify({"error": "No pages found in Figma file"}), 500
        page_id = pages[0]["id"]
    except Exception as e:
        return _figma_push_via_plugin_payload(figma_spec, comp_name)
    node = _build_figma_node(figma_spec, comp_name)
    try:
        create_resp = http_requests.post(f"{FIGMA_API_BASE}/files/{FIGMA_FILE_KEY}/nodes",
            headers=headers, json={"nodes": [node], "parent": {"type": "PAGE", "id": page_id}}, timeout=20)
        if create_resp.status_code in [200, 201]:
            result = create_resp.json(); node_id = result.get("nodes", [{}])[0].get("id", "")
            figma_url = f"https://www.figma.com/file/{FIGMA_FILE_KEY}?node-id={node_id}"
            return jsonify({"ok": True, "node_id": node_id, "figma_url": figma_url, "message": f"'{comp_name}' pushed to Figma successfully"})
        else:
            return _figma_push_via_plugin_payload(figma_spec, comp_name)
    except Exception as e:
        return _figma_push_via_plugin_payload(figma_spec, comp_name)

def _build_figma_node(spec: dict, name: str) -> dict:
    fills   = spec.get("fills",   [{"type": "SOLID", "r": 1.0, "g": 1.0, "b": 1.0, "a": 1.0}])
    strokes = spec.get("strokes", [{"type": "SOLID", "r": 0.878, "g": 0.878, "b": 0.878, "a": 1.0}])
    return {"type":"FRAME","name":name,"width":spec.get("width",400),"height":spec.get("height",200),"layoutMode":spec.get("layoutMode","VERTICAL"),"primaryAxisSizingMode":spec.get("primaryAxisSizingMode","AUTO"),"counterAxisSizingMode":spec.get("counterAxisSizingMode","FIXED"),"paddingTop":spec.get("paddingTop",24),"paddingRight":spec.get("paddingRight",24),"paddingBottom":spec.get("paddingBottom",24),"paddingLeft":spec.get("paddingLeft",24),"itemSpacing":spec.get("itemSpacing",16),"cornerRadius":spec.get("cornerRadius",0),"fills":fills,"strokes":strokes,"strokeWeight":spec.get("strokeWeight",1),"description":spec.get("description",f"Generated by Ai4UX — {name}"),"children":spec.get("children",[])}

def _figma_push_via_plugin_payload(spec: dict, name: str):
    plugin_script = f"""// Ai4UX — Auto-generated Figma plugin script
// Paste in Figma → Plugins → Development → Open Console
const spec = {json.dumps(spec, indent=2)};
async function createComponent() {{
  await figma.loadFontAsync({{ family: "Inter", style: "Regular" }});
  await figma.loadFontAsync({{ family: "Inter", style: "SemiBold" }});
  const frame = figma.createFrame();
  frame.name = "{name}";
  frame.resize(spec.width || 400, spec.height || 200);
  frame.layoutMode = spec.layoutMode || "VERTICAL";
  frame.paddingTop = spec.paddingTop || 24;
  frame.paddingRight = spec.paddingRight || 24;
  frame.paddingBottom = spec.paddingBottom || 24;
  frame.paddingLeft = spec.paddingLeft || 24;
  frame.itemSpacing = spec.itemSpacing || 16;
  frame.cornerRadius = spec.cornerRadius || 0;
  if (spec.fills && spec.fills[0]) {{
    const f = spec.fills[0];
    frame.fills = [{{ type: "SOLID", color: {{ r: f.r, g: f.g, b: f.b }}, opacity: f.a || 1 }}];
  }}
  for (const child of (spec.children || [])) {{
    if (child.type === "TEXT") {{
      const text = figma.createText();
      text.name = child.name || "Text";
      text.characters = child.characters || "";
      text.fontSize = child.fontSize || 14;
      frame.appendChild(text);
    }} else if (child.type === "FRAME" || child.type === "RECTANGLE") {{
      const rect = figma.createFrame();
      rect.name = child.name || "Frame";
      rect.resize(child.width || 100, child.height || 40);
      frame.appendChild(rect);
    }}
  }}
  figma.currentPage.appendChild(frame);
  figma.viewport.scrollAndZoomIntoView([frame]);
  figma.notify("✅ {name} created by Ai4UX!");
}}
createComponent().catch(console.error);""".strip()
    return jsonify({"ok":True,"method":"plugin_script","component_name":name,"plugin_script":plugin_script,"instructions":["1. Open your Figma file","2. Go to Menu → Plugins → Development → Open Console","3. Paste the plugin_script","4. Press Enter"],"figma_url":f"https://www.figma.com/file/{FIGMA_FILE_KEY}"})

@app.route("/api/figma/status")
@login_required
def figma_status():
    if not FIGMA_TOKEN or not FIGMA_FILE_KEY:
        return jsonify({"configured": False, "message": "Set FIGMA_TOKEN and FIGMA_FILE_KEY in Railway"})
    try:
        resp = http_requests.get(f"{FIGMA_API_BASE}/files/{FIGMA_FILE_KEY}?depth=1",
            headers={"X-Figma-Token": FIGMA_TOKEN}, timeout=10)
        if resp.status_code == 200:
            file_name = resp.json().get("name", "Unknown file")
            return jsonify({"configured": True, "file_name": file_name, "file_key": FIGMA_FILE_KEY, "figma_url": f"https://www.figma.com/file/{FIGMA_FILE_KEY}"})
        elif resp.status_code == 403:
            return jsonify({"configured": False, "message": "Token invalid or no access to file"})
        else:
            return jsonify({"configured": False, "message": f"Figma API returned {resp.status_code}"})
    except Exception as e:
        return jsonify({"configured": False, "message": str(e)})

# ── Entry point ───────────────────────────────────────────────
if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    debug = not os.environ.get("RAILWAY_ENVIRONMENT")
    app.run(host="0.0.0.0", port=port, debug=debug)

    # ============================================================
# SPRINT 8 — Jira Ticket Intelligence
# Add this route to ai4ux.py before the if __name__ block
# ============================================================

TICKET_INTELLIGENCE_PROMPT = """You are an AI Product Requirement Intelligence Assistant integrated with Jira.
Your job is to transform vague or incomplete Jira tickets into implementation-ready product requirements and UX acceptance criteria.

You will receive a Jira ticket with title, description, acceptance criteria, labels, and status.

Your responsibilities:
1. Understand the product intent behind the ticket.
2. Detect ambiguity, missing requirements, missing edge cases, and unclear UX behavior.
3. Perform lightweight competitive and industry-standard UX analysis using your knowledge of common SaaS/product patterns.
4. Generate structured implementation-ready output.

IMPORTANT RULES:
- Do NOT generate generic boilerplate.
- Think like a senior Product Manager + UX strategist.
- Focus on usability, edge cases, accessibility, enterprise SaaS workflows, and implementation clarity.
- Infer missing requirements carefully but clearly label assumptions.
- Keep outputs concise, structured, and implementation-friendly.
- Prioritize practical workflows over theoretical suggestions.
- Acceptance criteria should be testable.
- UX criteria should help both designers and developers.

Return ONLY valid JSON with this exact structure:
{
  "ticket_understanding": "2-3 sentence summary of the intended feature/problem",
  "competitive_insights": [
    {"pattern": "Pattern name", "description": "How leading SaaS products handle this", "examples": ["Product A", "Product B"]}
  ],
  "clarification_questions": [
    {"question": "Question text", "why_it_matters": "Impact on implementation"}
  ],
  "enhanced_description": {
    "objective": "What this feature achieves",
    "user_problem": "The specific problem being solved",
    "proposed_experience": "How the experience should work",
    "scope": "What is and is not included",
    "dependencies": "Technical or product dependencies",
    "assumptions": "Clearly labeled assumptions made"
  },
  "ux_acceptance_criteria": [
    {"category": "Loading States|Empty States|Error States|Accessibility|Responsiveness|Validation|Edge Cases", "criteria": "Testable criteria item", "priority": "Must Have|Should Have|Nice to Have"}
  ],
  "edge_cases": [
    {"case": "Edge case description", "risk": "High|Medium|Low", "recommendation": "How to handle it"}
  ],
  "suggested_priority": "High|Medium|Low",
  "priority_rationale": "One sentence explaining the priority"
}

Tone: Professional, concise, enterprise-product focused, actionable.
Avoid: excessive verbosity, generic AI wording, filler text."""


@app.route("/enrich-ticket", methods=["POST"])
@login_required
def enrich_ticket():
    """
    Enrich a Jira ticket with PM/UX intelligence.
    Accepts: { ticket_id: "PROJ-123" } or { ticket: {...} } (pre-fetched)
    Returns: structured intelligence JSON
    """
    data = request.get_json()
    ticket_id = data.get("ticket_id", "").strip()
    ticket    = data.get("ticket", None)

    # Fetch ticket if not provided
    if not ticket and ticket_id:
        try:
            ticket = fetch_jira_ticket(ticket_id)
        except Exception as e:
            return jsonify({"error": str(e)}), 400

    if not ticket:
        return jsonify({"error": "No ticket provided"}), 400

    uid = current_user_id()

    # Build ticket context string
    ticket_context = f"""JIRA TICKET:
Key: {ticket.get('key', '')}
Summary: {ticket.get('summary', '')}
Type: {ticket.get('issue_type', '')}
Status: {ticket.get('status', '')}
Priority: {ticket.get('priority', '')}
Assignee: {ticket.get('assignee', '')}
Labels: {', '.join(ticket.get('labels', [])) or 'None'}

Description:
{ticket.get('description', 'No description provided')}

Acceptance Criteria:
{ticket.get('acceptance_criteria', 'No acceptance criteria defined')}"""

    # Pull RAG context for additional intelligence
    try:
        rag_ctx = rag.build_analysis_context(
            components_detected=[],
            user_id=uid,
            screen_description=ticket.get('summary', '') + ' ' + ticket.get('description', '')[:300]
        )
    except:
        rag_ctx = ""

    # Add conventions context
    conv_ctx = get_conventions_context(uid)
    prod_ctx = get_product_context_summary(uid)

    full_context = ticket_context
    if rag_ctx:   full_context += f"\n\nORGANISATIONAL CONTEXT:\n{rag_ctx}"
    if conv_ctx:  full_context += f"\n\nESTABLISHED CONVENTIONS:\n{conv_ctx}"
    if prod_ctx:  full_context += f"\n\nPRODUCT CONTEXT:\n{prod_ctx}"

    try:
        msg = client.messages.create(
            model="claude-sonnet-4-6",
            max_tokens=4000,
            system=TICKET_INTELLIGENCE_PROMPT,
            messages=[{
                "role": "user",
                "content": f"Analyse this Jira ticket and return the structured intelligence JSON:\n\n{full_context}"
            }]
        )
        raw = msg.content[0].text.strip()
        if raw.startswith("```"):
            raw = raw.split("```")[1]
            raw = raw[4:].strip() if raw.startswith("json") else raw.strip()
        result = json.loads(raw)
        result["_ticket"] = ticket
        result["_enriched_at"] = datetime.now().strftime("%Y-%m-%d %H:%M")
        return jsonify(result)
    except json.JSONDecodeError as e:
        return jsonify({"error": f"Could not parse AI response: {str(e)}"}), 500
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/enrich-ticket/pdf", methods=["POST"])
@login_required
def enrich_ticket_pdf():
    """Generate a PDF from enriched ticket intelligence."""
    data = request.get_json()
    result = data.get("result", {})
    ticket = result.get("_ticket", {})

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter,
        leftMargin=inch, rightMargin=inch, topMargin=inch, bottomMargin=inch)
    S    = getSampleStyleSheet()
    BLUE = colors.HexColor("#3A6FF7")
    DARK = colors.HexColor("#0D1117")
    GREY = colors.HexColor("#8A95A3")
    LGREY= colors.HexColor("#f4f4f4")
    GREEN= colors.HexColor("#198038")
    RED  = colors.HexColor("#da1e28")
    GOLD = colors.HexColor("#b28600")

    ts  = ParagraphStyle('T', parent=S['Normal'], fontSize=20, textColor=DARK, fontName='Helvetica-Bold', spaceAfter=4)
    ss  = ParagraphStyle('S', parent=S['Normal'], fontSize=10, textColor=GREY, spaceAfter=16, fontName='Helvetica')
    h1  = ParagraphStyle('H1', parent=S['Normal'], fontSize=12, textColor=BLUE, fontName='Helvetica-Bold', spaceBefore=18, spaceAfter=6)
    bs  = ParagraphStyle('B', parent=S['Normal'], fontSize=10, leading=15, textColor=DARK, spaceAfter=4)
    bl  = ParagraphStyle('BL', parent=S['Normal'], fontSize=10, leading=15, textColor=DARK, leftIndent=16, spaceAfter=3)
    sm  = ParagraphStyle('SM', parent=S['Normal'], fontSize=9, textColor=GREY, spaceAfter=2)

    story = []

    # Header
    story.append(Paragraph("Aetheris — Ticket Intelligence Report", ts))
    story.append(Paragraph(
        f"Ticket: {ticket.get('key','')}  ·  {result.get('_enriched_at','')}  ·  Priority: {result.get('suggested_priority','')}",
        ss))

    # Ticket summary
    story.append(Paragraph("TICKET", h1))
    tbl = Table([
        ["Key",      ticket.get('key','')],
        ["Summary",  ticket.get('summary','')],
        ["Type",     ticket.get('issue_type','')],
        ["Status",   ticket.get('status','')],
        ["Priority", ticket.get('priority','')],
    ], colWidths=[1.2*inch, 5.3*inch])
    tbl.setStyle(TableStyle([
        ('FONTNAME',(0,0),(0,-1),'Helvetica-Bold'),('FONTSIZE',(0,0),(-1,-1),9),
        ('TEXTCOLOR',(0,0),(0,-1),BLUE),
        ('ROWBACKGROUNDS',(0,0),(-1,-1),[colors.white,LGREY]),
        ('GRID',(0,0),(-1,-1),0.5,colors.HexColor("#e0e0e0")),
        ('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5),
        ('LEFTPADDING',(0,0),(-1,-1),8)
    ]))
    story.append(tbl); story.append(Spacer(1,8))

    # Ticket understanding
    story.append(Paragraph("TICKET UNDERSTANDING", h1))
    story.append(Paragraph(result.get('ticket_understanding',''), bs))

    # Enhanced description
    enh = result.get('enhanced_description', {})
    if enh:
        story.append(Paragraph("ENHANCED STORY DESCRIPTION", h1))
        for field, label in [
            ('objective','Objective'), ('user_problem','User Problem'),
            ('proposed_experience','Proposed Experience'), ('scope','Scope'),
            ('dependencies','Dependencies'), ('assumptions','Assumptions')
        ]:
            if enh.get(field):
                story.append(Paragraph(f"<b>{label}:</b>", bs))
                story.append(Paragraph(enh[field], bl))

    # UX Acceptance Criteria
    criteria = result.get('ux_acceptance_criteria', [])
    if criteria:
        story.append(Paragraph("UX ACCEPTANCE CRITERIA", h1))
        td = [["Category", "Criteria", "Priority"]]
        for c in criteria:
            pri_color = GREEN if c.get('priority')=='Must Have' else GOLD if c.get('priority')=='Should Have' else GREY
            td.append([
                Paragraph(c.get('category',''), sm),
                Paragraph(c.get('criteria',''), bs),
                Paragraph(f'<font color="#{pri_color.hexval()[2:]}"><b>{c.get("priority","")}</b></font>', sm)
            ])
        t = Table(td, colWidths=[1.2*inch, 4.0*inch, 1.3*inch], repeatRows=1)
        t.setStyle(TableStyle([
            ('BACKGROUND',(0,0),(-1,0),DARK),('TEXTCOLOR',(0,0),(-1,0),colors.white),
            ('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),('FONTSIZE',(0,0),(-1,-1),9),
            ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,LGREY]),
            ('GRID',(0,0),(-1,-1),0.5,colors.HexColor("#e0e0e0")),
            ('VALIGN',(0,0),(-1,-1),'TOP'),
            ('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5),
            ('LEFTPADDING',(0,0),(-1,-1),6)
        ]))
        story.append(t); story.append(Spacer(1,8))

    # Edge cases
    edges = result.get('edge_cases', [])
    if edges:
        story.append(Paragraph("EDGE CASES & RISKS", h1))
        for e in edges:
            risk_color = RED if e.get('risk')=='High' else GOLD if e.get('risk')=='Medium' else GREY
            story.append(Paragraph(
                f'<font color="#{risk_color.hexval()[2:]}"><b>[{e.get("risk","")}]</b></font> {e.get("case","")}', bs))
            story.append(Paragraph(f"→ {e.get('recommendation','')}", bl))

    # Clarification questions
    questions = result.get('clarification_questions', [])
    if questions:
        story.append(Paragraph("CLARIFICATION QUESTIONS", h1))
        for i, q in enumerate(questions, 1):
            story.append(Paragraph(f"<b>Q{i}:</b> {q.get('question','')}", bs))
            story.append(Paragraph(f"Why it matters: {q.get('why_it_matters','')}", bl))

    # Competitive insights
    insights = result.get('competitive_insights', [])
    if insights:
        story.append(Paragraph("COMPETITIVE & BEST-PRACTICE INSIGHTS", h1))
        for ins in insights:
            story.append(Paragraph(f"<b>{ins.get('pattern','')}</b>", bs))
            story.append(Paragraph(ins.get('description',''), bl))
            if ins.get('examples'):
                story.append(Paragraph(f"Examples: {', '.join(ins['examples'])}", sm))

    story.append(Spacer(1,20))
    story.append(Paragraph("Generated by Aetheris — Design Intelligence",
        ParagraphStyle('F', parent=S['Normal'], fontSize=8, textColor=GREY, alignment=1)))

    doc.build(story); buf.seek(0)
    ticket_key = ticket.get('key','ticket').replace('-','_').lower()
    return send_file(buf, mimetype="application/pdf", as_attachment=True,
        download_name=f"aetheris_intelligence_{ticket_key}.pdf")

